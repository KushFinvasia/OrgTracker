"""The monthly check-in workbook: one file in which a Manager marks the whole team, then uploads it.

Sheets: Check-ins (ratings, training score, comments), KPIs (achieved), Duty tasks (done?, completed on, delay reason),
Training (module completion dates) and How to use. Rows are matched on upload by Employee ID plus the KPI, task or
module ID in the hidden "Row ID" column; the hidden "Version" column lets the upload spot rows changed in Finvasia
Elevate after the file was downloaded. Building and reading live here; matching and the business rules stay in the
route, which applies every change through the same helpers as the person page.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

# Shoonya tokens (CLAUDE.md section 17), as ARGB without the leading #.
INK, MUTED, GREEN, MINT, BORDER = "2D2D2D", "69766F", "23763F", "E7F5EB", "E2EAE4"
EDIT_FILL, LOCKED_FILL, READONLY_FILL, AMBER_FILL = "FFF6DD", "F6FAF7", "ECEEEC", "FCF3D9"

SHEET_CHECKINS, SHEET_KPIS, SHEET_DUTIES, SHEET_TRAINING, SHEET_HELP = "Check-ins", "KPIs", "Duty tasks", "Training", "How to use"
HEADER_ROW = 4  # rows 1-3 hold the title, the scope line and the legend
ROW_ID, VERSION = "Row ID", "Version"

# Indian digit grouping (12,34,567) for counts and rupees; Excel picks the section by size.
NUM_FMT = r'[>=10000000]##\,##\,##\,##0;[>=100000]##\,##\,##0;##,##0'
RUPEE_FMT = r'[>=10000000]"₹ "##\,##\,##\,##0;[>=100000]"₹ "##\,##\,##0;"₹ "##,##0'
PCT_UNIT_FMT = '0.00"%"'
DATE_FMT = "dd mmm yyyy"


@dataclass
class Col:
    header: str
    key: str
    width: float
    kind: str = "text"  # text | number | date | rating | twcs | outcome | pct | result | hidden
    edit: bool = False
    wrap: bool = False


@dataclass
class Row:
    values: dict
    editable: bool
    note: str | None = None  # why a row is read-only, shown as a cell comment
    number_format: str | None = None  # for the KPI target and achieved cells
    formula_extra: dict = field(default_factory=dict)


@dataclass
class Book:
    month_label: str
    scope_line: str
    rating_cols: list[tuple[str, str]]  # (header, field)
    checkins: list[Row]
    kpis: list[Row]
    duties: list[Row]
    training: list[Row]
    today: date
    month_closed: bool = False


def value_format(unit: str | None) -> str:
    return RUPEE_FMT if unit == "rupees" else PCT_UNIT_FMT if unit == "percent" else NUM_FMT


def _columns(book: Book) -> dict[str, list[Col]]:
    ratings = [Col(h, f, 11, "rating", edit=True) for h, f in book.rating_cols]
    return {
        SHEET_CHECKINS: [
            Col("Employee ID", "code", 13), Col("Name", "name", 24), Col("Designation", "designation", 24), Col("Manager", "manager", 22),
            Col("Check-in", "state", 14), *ratings, Col("Training (TWCS)", "twcs_score", 15, "twcs", edit=True),
            Col("Comments", "comments", 44, edit=True, wrap=True), Col(VERSION, "version", 8, "hidden"),
        ],
        SHEET_KPIS: [
            Col("Employee ID", "code", 13), Col("Name", "name", 22), Col("KPI", "title", 42, wrap=True), Col("KPI type", "type", 20),
            Col("Target", "target", 14, "number"), Col("Achieved", "achieved", 14, "number", edit=True), Col("% of target", "pct", 12, "pct"),
            Col("Result", "result", 16, "result"), Col(ROW_ID, "row_id", 8, "hidden"), Col(VERSION, "version", 8, "hidden"),
        ],
        SHEET_DUTIES: [
            Col("Employee ID", "code", 13), Col("Name", "name", 22), Col("Task", "title", 44, wrap=True), Col("Due date", "due", 13, "date"),
            Col("Done?", "outcome", 16, "outcome", edit=True), Col("Completed on", "completed", 15, "date", edit=True),
            Col("Delay reason", "delay_reason", 40, edit=True, wrap=True), Col(ROW_ID, "row_id", 8, "hidden"), Col(VERSION, "version", 8, "hidden"),
        ],
        SHEET_TRAINING: [
            Col("Employee ID", "code", 13), Col("Name", "name", 22), Col("Module", "title", 40, wrap=True), Col("Track", "track", 14),
            Col("Due date", "due", 13, "date"), Col("Completed on", "completed", 15, "date", edit=True),
            Col(ROW_ID, "row_id", 8, "hidden"), Col(VERSION, "version", 14, "hidden"),
        ],
    }


PURPOSE = {
    SHEET_CHECKINS: "Ratings, training score and comments, one row per person",
    SHEET_KPIS: "Achieved value for each KPI. % of target and Result work themselves out",
    SHEET_DUTIES: "Whether each duty task was done, when, and why it was late",
    SHEET_TRAINING: "Completion dates for the training modules due this month",
}


def build(book: Book) -> bytes:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    fill = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731
    thin = Side(style="thin", color=BORDER)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    excel_date = lambda d: f"DATE({d.year},{d.month},{d.day})"  # noqa: E731

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cols_by_sheet = _columns(book)
    rows_by_sheet = {SHEET_CHECKINS: book.checkins, SHEET_KPIS: book.kpis, SHEET_DUTIES: book.duties, SHEET_TRAINING: book.training}

    for name, cols in cols_by_sheet.items():
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        last = get_column_letter(len(cols))
        ws["A1"] = f"Finvasia Elevate · Monthly check-in · {book.month_label}"
        ws["A1"].font = Font(bold=True, size=14, color=INK)
        ws["A2"] = f"{PURPOSE[name]}. {book.scope_line}"
        ws["A2"].font = Font(size=10, color=MUTED)
        ws["A3"] = ("This month is closed: the file is read-only." if book.month_closed
                    else "Edit the tinted cells only. Grey rows are read-only (submitted, approved or closed). Blank cells leave values unchanged.")
        ws["A3"].font = Font(size=10, italic=True, color=GREEN if not book.month_closed else MUTED)
        for r in (1, 2, 3):
            ws.merge_cells(f"A{r}:{last}{r}")
        ws.row_dimensions[1].height = 24

        for i, c in enumerate(cols, start=1):
            cell = ws.cell(HEADER_ROW, i, c.header)
            cell.font = Font(bold=True, color=INK)
            cell.fill = fill(MINT)
            cell.border = Border(bottom=Side(style="medium", color=GREEN), left=thin, right=thin, top=thin)
            cell.alignment = Alignment(vertical="center", horizontal="right" if c.kind in ("number", "pct", "twcs", "rating") else "left", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = c.width
            if c.kind == "hidden":
                ws.column_dimensions[get_column_letter(i)].hidden = True
        ws.row_dimensions[HEADER_ROW].height = 30

        validations: dict[str, DataValidation] = {}

        def dv(kind: str) -> DataValidation:
            if kind not in validations:
                if kind == "rating":
                    v = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True, error="Choose a whole number from 1 to 5.", errorTitle="Rating")
                elif kind == "outcome":
                    v = DataValidation(type="list", formula1='"Completed,Not completed"', allow_blank=True, error="Choose Completed or Not completed.", errorTitle="Done?")
                elif kind == "twcs":
                    v = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True, error="Enter a score from 0 to 100.", errorTitle="Training score")
                elif kind == "number":
                    v = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, error="Enter a number of 0 or more.", errorTitle="Achieved")
                else:  # date
                    v = DataValidation(type="date", operator="between", formula1=excel_date(date(2020, 1, 1)), formula2=excel_date(book.today), allow_blank=True,
                                       error=f"Enter a date no later than today ({book.today:%d %b %Y}), e.g. {book.today:%d %b %Y}.", errorTitle="Date")
                v.showErrorMessage = True
                ws.add_data_validation(v)
                validations[kind] = v
            return validations[kind]

        col_index = {c.key: i for i, c in enumerate(cols, start=1)}
        for n, row in enumerate(rows_by_sheet[name], start=HEADER_ROW + 1):
            editable = row.editable and not book.month_closed
            for i, c in enumerate(cols, start=1):
                letter = get_column_letter(i)
                cell = ws.cell(n, i)
                v = row.values.get(c.key)
                if c.kind == "pct":
                    t, a = f"{get_column_letter(col_index['target'])}{n}", f"{get_column_letter(col_index['achieved'])}{n}"
                    fallback = row.formula_extra.get("achieved_fallback")
                    num = f"IF(ISNUMBER({a}),{a},{fallback})" if fallback is not None else a
                    cond = f"AND(ISNUMBER({t}),{t}>0)" if fallback is not None else f"AND(ISNUMBER({t}),{t}>0,ISNUMBER({a}))"
                    cell.value = f'=IF({cond},{num}/{t},"")'
                    cell.number_format = "0.00%"
                elif c.kind == "result":
                    p = f"{get_column_letter(col_index['pct'])}{n}"
                    cell.value = f'=IF({p}="","",IF({p}<1,"Under achieved",IF({p}>1,"Over achieved","Achieved")))'
                else:
                    cell.value = v
                if c.kind == "date":
                    cell.number_format = DATE_FMT
                elif c.kind == "number":
                    cell.number_format = row.number_format or NUM_FMT
                elif c.kind == "twcs":
                    cell.number_format = "0.##"
                cell.border = box
                cell.alignment = Alignment(vertical="top", wrap_text=c.wrap, horizontal="right" if c.kind in ("number", "pct", "twcs", "rating") else None)
                can = c.edit and editable
                if can:
                    cell.fill = fill(EDIT_FILL)
                    cell.protection = Protection(locked=False)
                    if c.kind in ("rating", "outcome", "twcs", "number", "date"):
                        dv(c.kind).add(f"{letter}{n}")
                elif not editable:
                    cell.fill = fill(READONLY_FILL)
                    cell.font = Font(color=MUTED)
                else:
                    cell.fill = fill(LOCKED_FILL)
            if row.note:
                ws.cell(n, 1).comment = Comment(row.note, "Finvasia Elevate")

        if name == SHEET_KPIS and rows_by_sheet[name]:
            rng = f"{get_column_letter(col_index['result'])}{HEADER_ROW + 1}:{get_column_letter(col_index['result'])}{HEADER_ROW + len(rows_by_sheet[name])}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Under achieved"'], fill=fill(AMBER_FILL), font=Font(color="85650F", bold=True)))
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Achieved"'], fill=fill(MINT), font=Font(color=GREEN, bold=True)))
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Over achieved"'], fill=fill(MINT), font=Font(color=GREEN, bold=True)))
        if not rows_by_sheet[name]:
            ws.cell(HEADER_ROW + 1, 1, "Nothing to mark here this month.").font = Font(italic=True, color=MUTED)

        ws.freeze_panes = ws.cell(HEADER_ROW + 1, 3)  # header row plus Employee ID and Name
        ws.auto_filter.ref = f"A{HEADER_ROW}:{last}{HEADER_ROW + max(len(rows_by_sheet[name]), 1)}"
        ws.protection.sheet = True  # no password: the server enforces every rule; this guards against accidental edits
        ws.protection.autoFilter = False
        ws.protection.sort = False
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"

    _help_sheet(wb.create_sheet(SHEET_HELP), book)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _help_sheet(ws, book: Book) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    ws["A1"] = f"Finvasia Elevate · Monthly check-in · {book.month_label}"
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    ws["A2"] = book.scope_line
    ws["A2"].font = Font(size=10, color=MUTED)
    ratings = " and ".join(h for h, _ in book.rating_cols) or "The ratings"
    lines = [
        ("Steps", "1. Mark the tinted cells on each sheet. 2. Save the file as .xlsx. 3. On the Team page choose Upload marked template, check the preview, then apply."),
        ("What you can edit", f"Check-ins: {ratings}, Training (TWCS) and Comments. KPIs: Achieved. Duty tasks: Done?, Completed on and Delay reason. "
                              "Training: Completed on. Everything else is read-only and is worked out in Finvasia Elevate."),
        ("Blank cells", "A blank cell leaves the value unchanged. To clear a value, use the person's page in Team."),
        ("Ratings", f"{ratings}: whole numbers from 1 (well below expectations) to 5 (well above). Only the ratings weighted for the team's category are asked for."),
        ("Training (TWCS)", "Optional, 0 to 100. Leave it blank and the score is worked out from the module completion dates on the Training sheet."),
        ("KPIs", "Type the achieved value for the month. % of target and Result (Under achieved below 100%, Achieved at 100%, Over achieved above) "
                 "fill in by themselves. KPIs have no completion date or delay: On-time comes from duty tasks."),
        ("Duty tasks", "Choose Completed or Not completed. A completed task needs the date it was done. Add a delay reason when it was late; "
                       "the HOD or an Admin validates it in Finvasia Elevate."),
        ("Dates", f"Type a date such as {book.today:%d %b %Y} (or {book.today:%d/%m/%Y}). Dates cannot be in the future."),
        ("Read-only rows", "Grey rows cannot change: the check-in is submitted or approved, the month is closed, or someone else marks that person. "
                           "A returned check-in can be marked again."),
        ("Matching", "Rows are matched by Employee ID and a hidden Row ID. Don't delete, insert or reorder columns. You can sort and filter rows."),
        ("Changed since download", "If someone changed a row in Finvasia Elevate after you downloaded this file, that row is not applied. Download a fresh file."),
        ("Audit", "Every change from the file is recorded in the person's history with the file name, sheet and row."),
    ]
    for i, (a, b) in enumerate(lines, start=4):
        ws.cell(i, 1, a).font = Font(bold=True, color=INK)
        ws.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 1).alignment = Alignment(vertical="top")
        if i % 2 == 0:
            for c in (1, 2):
                ws.cell(i, c).fill = PatternFill("solid", fgColor="F6FAF7")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# --- reading ---------------------------------------------------------------------------------------

HEADER_KEYS = {
    SHEET_KPIS: {"employee id": "code", "name": "name", "kpi": "title", "achieved": "achieved", "done?": "outcome", "completed on": "completed",
                 "delay reason": "delay_reason", "row id": "row_id", "version": "version"},
    SHEET_DUTIES: {"employee id": "code", "name": "name", "task": "title", "done?": "outcome", "completed on": "completed",
                   "delay reason": "delay_reason", "row id": "row_id", "version": "version"},
    SHEET_TRAINING: {"employee id": "code", "name": "name", "module": "title", "completed on": "completed", "row id": "row_id", "version": "version"},
}


def _norm(h) -> str:
    return " ".join(str(h or "").replace("\n", " ").lower().split())


@dataclass
class SheetRows:
    name: str
    columns: set[str]
    rows: list[tuple[int, dict]]  # (Excel row number, {key: raw value})


def read_sheet(ws, keys: dict[str, str]) -> SheetRows | None:
    """Rows under the first header row that starts with Employee ID. Exact header names, so the hidden columns are found too."""
    it = ws.iter_rows(values_only=True)
    cols: dict[int, str] | None = None
    row_no = 0
    for row in it:
        row_no += 1
        found = {i: keys[_norm(h)] for i, h in enumerate(row) if _norm(h) in keys}
        if "code" in found.values() and len(found) >= 2:
            cols = found
            break
    if cols is None:
        return None
    out = []
    for row in it:
        row_no += 1
        vals = {k: row[i] if i < len(row) else None for i, k in cols.items()}
        if all(blank(v) for v in vals.values()):
            continue
        out.append((row_no, vals))
    return SheetRows(ws.title, set(cols.values()), out)


def blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def text(v) -> str | None:
    return None if blank(v) else " ".join(str(v).split())


def parse_number(v) -> tuple[float | None, str | None]:
    if blank(v):
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None
    s = str(v).strip().replace(",", "").replace("₹", "").replace("Rs.", "").replace("Rs", "").strip()
    try:
        return float(s), None
    except ValueError:
        return None, f"'{v}' is not a number"


_DATE_FORMATS = ("%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %y")


def parse_date(v) -> tuple[date | None, str | None]:
    if blank(v):
        return None, None
    if isinstance(v, datetime):
        return v.date(), None
    if isinstance(v, date):
        return v, None
    s = re.sub(r"\s+", " ", str(v).strip())
    for f in _DATE_FORMATS:
        try:
            return datetime.strptime(s, f).date(), None
        except ValueError:
            continue
    return None, f"'{v}' is not a date. Use a date such as 09 Oct 2026"


def parse_outcome(v) -> tuple[str | None, str | None]:
    if blank(v):
        return None, None
    s = _norm(v)
    if s in ("completed", "complete", "done", "yes"):
        return "Completed", None
    if s in ("not completed", "not complete", "not done", "no"):
        return "Not completed", None
    return None, f"'{v}' must be Completed or Not completed"


def parse_row_id(v) -> int | None:
    n, err = parse_number(v)
    return int(n) if n is not None and not err and n == int(n) else None
