"""Individual Tier (Sales pilot) API: months, KPI rows, duty tasks, TWCS training, evaluations, lock, rewards, reports, KPI Scoring.

Every list is limited to the viewer's pilot scope on the server (`services.individual.visible_ids`).
"""

from __future__ import annotations

import copy
import csv
import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select

from app.api.deps import Ctx, get_ctx
from app.core.errors import Conflict, Forbidden, NotFound, RuleViolation
from app.models import Account, AuditEvent, CorrectionRequest, Department, MonthlyNomination, ImportBatch, KpiItem, KpiMonth, KpiType, MonthlyEvaluation, ScorecardSnapshot, ScoringRule, TrainingCatalogItem, TrainingModule
from app.models.individual import NOMINATION_AWARDS
from app.seed.individual_demo import clear_individual_demo, has_demo, seed_individual_demo
from app.seed.sales import carry_over, ensure_month, month_end, open_due_months, training_cutoff
from app.services import individual as ind
from app.services import sales_scoring as ss
from app.services.audit import notify, record
from app.services.common import account_person, person, today
from app.core.clock import now_ist

router = APIRouter(prefix="/individual", tags=["individual"])


def _pilot(ctx: Ctx) -> ind.Pilot:
    p = ind.load_pilot(ctx.db)
    if p is None:
        raise NotFound("Sales pilot")
    return p


def _types(ctx: Ctx) -> dict[int, KpiType]:
    return {t.id: t for t in ctx.db.scalars(select(KpiType).order_by(KpiType.sort))}


def _require_open(month: KpiMonth, emp_id: int | None = None, db=None) -> None:
    """Closed months are frozen, except for one person whose correction an Admin has approved."""
    if month.state == "Locked" and not (emp_id is not None and db is not None and ind.open_correction(db, month, emp_id)):
        raise RuleViolation("The month is closed. Locked scores are frozen; changes need an approved correction.", "month_id")


def _month_dict(ctx: Ctx, m: KpiMonth, pilot: ind.Pilot | None = None) -> dict:
    rule = ctx.db.get(ScoringRule, m.rule_id)
    out = {
        "id": m.id,
        "month": m.month,
        "label": ind.month_label(m.month),
        "state": m.state,
        "training_cutoff": m.training_cutoff,
        "rule": {"id": rule.id, "label": rule.label, "state": rule.state},
        "locked_at": m.locked_at,
        "locked_by": account_person(ctx.db, m.locked_by),
        "lock_note": m.lock_note,
        "published_at": m.published_at,
        "illustrative": m.illustrative,
        "published_by": account_person(ctx.db, m.published_by),
    }
    if pilot is not None:
        ids = ind.visible_ids(pilot, ctx.scope)
        evs = list(ctx.db.scalars(select(MonthlyEvaluation).where(MonthlyEvaluation.month_id == m.id, MonthlyEvaluation.employee_id.in_(ids))))
        out["counts"] = {
            "people": len(ids),
            "approved": sum(e.state == "Approved" for e in evs),
            "submitted": sum(e.state == "Submitted" for e in evs),
            "returned": sum(e.state == "Returned" for e in evs),
            "rows": len(list(ctx.db.scalars(select(KpiItem.id).where(KpiItem.month_id == m.id, KpiItem.employee_id.in_(ids), KpiItem.removed_at.is_(None))))),
        }
    return out


# --- context -------------------------------------------------------------------------------------


@router.get("/context")
def context(ctx: Ctx = Depends(get_ctx)):
    """Who the viewer is in the pilot and which months exist. Opens the current month (and carries targets) when it is due."""
    pilot = ind.load_pilot(ctx.db)
    if pilot is None:
        return {"in_pilot": False, "months": [], "people": 0}
    if open_due_months(ctx.db):
        ctx.db.commit()
    ids = ind.visible_ids(pilot, ctx.scope)
    evaluates = [e.id for e in pilot.people if (ev := pilot.evaluator(e)) is not None and ev.id == ctx.account.id]
    months = list(ctx.db.scalars(select(KpiMonth).order_by(KpiMonth.month.desc())))
    current = next((m for m in months if m.state == "Open" and m.month <= today()), months[0] if months else None)
    # The earliest open month with work waiting is the one to act on first.
    oldest_open = next((m for m in reversed(months) if m.state == "Open"), None)
    return {
        "in_pilot": bool(ids) or ctx.account.role == "admin",
        "department": pilot.dept.name,
        "category": pilot.dept.category,
        "people": len(ids),
        "evaluates": len(evaluates),
        "role": ctx.account.role,
        "months": [_month_dict(ctx, m) for m in months],
        "current_month_id": current.id if current else None,
        "focus_month_id": oldest_open.id if oldest_open else (current.id if current else None),
        "demo_loaded": has_demo(ctx.db),
    }


@router.get("/kpi-types")
def kpi_types(ctx: Ctx = Depends(get_ctx)):
    return [{"id": t.id, "code": t.code, "name": t.name, "unit": t.unit, "description": t.description} for t in _types(ctx).values()]


@router.get("/months")
def months(ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    return [_month_dict(ctx, m, pilot) for m in ctx.db.scalars(select(KpiMonth).order_by(KpiMonth.month.desc()))]


class MonthIn(BaseModel):
    month: date


@router.post("/months")
def open_month(body: MonthIn, ctx: Ctx = Depends(get_ctx)):
    ctx.scope.require_role("admin", "hod")
    pilot = _pilot(ctx)
    first = body.month.replace(day=1)
    if ctx.db.scalar(select(KpiMonth).where(KpiMonth.month == first)):
        raise RuleViolation(f"{ind.month_label(first)} is already open.", "month")
    rule = _rule_for_month(ctx, first)
    m = ensure_month(ctx.db, first, rule, pilot.dept)
    carried = carry_over(ctx.db, m, ctx.account.id)
    record(ctx.db, ctx.account.id, "open_month", "kpi_month", m.id, ind.month_label(first), new={"rule": rule.label, "carried_over": carried})
    ctx.db.commit()
    return _month_dict(ctx, m, pilot)


def _rule_for_month(ctx: Ctx, first: date) -> ScoringRule:
    rule = ctx.db.scalar(
        select(ScoringRule).where(ScoringRule.state == "Approved", ScoringRule.effective_month <= first).order_by(ScoringRule.effective_month.desc(), ScoringRule.id.desc())
    )
    if rule is None:
        raise RuleViolation("No approved scoring rule covers this month.", "month")
    return rule


# --- people / team -------------------------------------------------------------------------------


def _people_rows(ctx: Ctx, pilot: ind.Pilot, month: KpiMonth) -> list[dict]:
    data = ind.load_month(pilot, month)
    ids = ind.visible_ids(pilot, ctx.scope)
    rows = []
    for emp in pilot.people:
        if emp.id not in ids:
            continue
        card = ind.compute(pilot, data, emp)
        ev = data.evaluations.get(emp.id)
        evaluator = pilot.evaluator(emp)
        ok, _ = ind.can_edit(pilot, ctx.account, emp, month, ev)
        flags = sum(ind.open_warnings(it) for it in data.items.get(emp.id, []))
        rows.append(
            {
                "employee": person(emp),
                "manager": person(emp.manager) if emp.manager else None,
                "evaluator": person(evaluator.employee) if evaluator else None,
                "evaluator_is_me": bool(evaluator and evaluator.id == ctx.account.id),
                "approver_roles": list(pilot.approver_roles(emp)),
                "evaluation_state": ev.state if ev else "Not started",
                "evaluation_decided_by": ev.decided_by if ev else None,
                "can_edit": ok,
                "can_approve": bool(ev and ev.state == "Submitted" and _can_approve(ctx, pilot, emp, ev)[0]),
                "kpi_rows": sum(1 for it in data.items.get(emp.id, []) if it.kind == "kpi"),
                "duty_rows": sum(1 for it in data.items.get(emp.id, []) if it.kind == "duty"),
                "achievement_pct": card["achievement"].get("raw"),
                # KPI average: how many KPIs with a numeric target have an achieved value (the average uses only those).
                "kpis_marked": len(card["achievement"].get("rows") or []),
                "kpis_total": len(card["achievement"].get("rows") or []) + sum(1 for x in card["achievement"].get("skipped") or [] if x["reason"] != "No numeric target"),
                "achievement_score": card["achievement"].get("score"),
                "on_time": card["on_time"],
                "twcs": card["training"]["score"],
                "twcs_status": card["training"]["status"],
                "band": card["training"]["band"],
                "keka": card["training"]["keka"],
                "coverage": card["scorecard"]["coverage"],
                "base": card["scorecard"]["base"],
                "final": card["scorecard"]["final"],
                "status": card["scorecard"]["status"] if not card.get("locked") else card.get("status", card["scorecard"]["status"]),
                "tier": card["scorecard"]["tier"],
                "penalty_total": card["scorecard"]["penalty_total"],
                "readiness": card["readiness"],
                "warnings": flags,
                "locked": card.get("locked", False),
                "check_in": {
                    **{f"{m}_rating": getattr(ev, f"{m}_rating") if ev else None for m in ss.RATED_METRICS},
                    "entered": (card.get("inputs") or {}).get("entered", {}),
                    "derived": (card.get("inputs") or {}).get("derived", {}),
                    "version": ev.version if ev else None,
                    "can_submit": ok and ev is not None,
                },
            }
        )
    if not ind.method_visible(ctx.account):
        for r in rows:
            for k in ("coverage", "base", "penalty_total", "achievement_score"):
                r.pop(k, None)
    return rows


@router.get("/months/{month_id}/people")
def month_people(month_id: int, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    weights = ctx.db.get(ScoringRule, month.rule_id).config["weights"].get(pilot.dept.category, {})
    return {"month": _month_dict(ctx, month, pilot), "rows": _people_rows(ctx, pilot, month),
            "required_ratings": [m for m in ss.RATED_METRICS if weights.get(m)], "method_visible": ind.method_visible(ctx.account)}


@router.get("/team")
def team(month_id: int | None = None, ctx: Ctx = Depends(get_ctx)):
    """Roster and hierarchy for the viewer's pilot scope, with month-to-date achievement."""
    pilot = _pilot(ctx)
    month = ctx.db.get(KpiMonth, month_id) if month_id else ctx.db.scalar(select(KpiMonth).where(KpiMonth.state == "Open").order_by(KpiMonth.month))
    rows = _people_rows(ctx, pilot, month) if month else []
    ids = ind.visible_ids(pilot, ctx.scope)
    visible_heads = [h for h in pilot.heads if ctx.scope.perf_people is None or ctx.account.employee_id == h.id]
    return {
        "month": _month_dict(ctx, month, pilot) if month else None,
        "department": pilot.dept.name,
        "heads": [person(h) for h in visible_heads],
        "rows": rows,
        "tree": [{"id": e.id, "manager_id": e.manager_id} for e in pilot.people if e.id in ids],
    }


# --- KPI rows and duty tasks ---------------------------------------------------------------------


@router.get("/months/{month_id}/items")
def month_items(month_id: int, employee_id: int | None = None, kind: str | None = None, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    ids = ind.visible_ids(pilot, ctx.scope)
    if employee_id is not None:
        ind.require_person(pilot, ctx.scope, employee_id)
        ids = {employee_id}
    data = ind.load_month(pilot, month)
    types = _types(ctx)
    by_emp = pilot.by_id()
    out = []
    for emp_id in sorted(ids, key=lambda i: by_emp[i].code):
        emp = by_emp[emp_id]
        ok, why = ind.can_edit(pilot, ctx.account, emp, month, data.evaluations.get(emp_id))
        for it in data.items.get(emp_id, []):
            if kind and it.kind != kind:
                continue
            out.append({**ind.item_dict(it, data.rule.config, data.sums, types), "employee": person(emp), "can_edit": ok, "locked_reason": None if ok else why,
                        "can_validate_delay": _can_validate_delay(ctx, it, month)})
    return out


class ItemIn(BaseModel):
    employee_id: int
    kind: str = Field(pattern="^(kpi|duty)$")
    title: str = Field(min_length=2, max_length=500)
    kpi_type_id: int | None = None
    is_team: bool = False
    target_value: float | None = Field(default=None, ge=0)
    target_date: date | None = None
    priority: str | None = Field(default=None, pattern="^(High|Medium|Low)$")


def _manager_due(d: date | None, month: KpiMonth) -> date | None:
    """A Manager-set due date must fall inside the row's month; the month's last day is stored as blank (the default)."""
    if d is None:
        return None
    if not month.month <= d <= month_end(month.month):
        raise RuleViolation(f"The due date must be within {ind.month_label(month.month)}.", "target_date")
    return None if d == month_end(month.month) else d


@router.post("/months/{month_id}/items")
def create_item(month_id: int, body: ItemIn, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    _require_open(month, body.employee_id, ctx.db)
    emp = ind.require_person(pilot, ctx.scope, body.employee_id)
    ind.require_edit(pilot, ctx.account, emp, month)
    if body.kind == "kpi" and not body.kpi_type_id:
        raise RuleViolation("Choose the KPI type.", "kpi_type_id")
    if body.kpi_type_id and not ctx.db.get(KpiType, body.kpi_type_id):
        raise RuleViolation("Unknown KPI type.", "kpi_type_id")
    it = KpiItem(
        month_id=month.id, employee_id=emp.id, kind=body.kind, title=body.title.strip(), kpi_type_id=body.kpi_type_id if body.kind == "kpi" else None,
        is_team=body.is_team and body.kind == "kpi", target_value=body.target_value if body.kind == "kpi" else None,
        achieved_value=0.0 if body.kind == "kpi" and not body.is_team else None,  # KPI achieved defaults to 0
        target_date=_manager_due(body.target_date, month), start_date=month.month, priority=body.priority,
        source="Entered in Finvasia Elevate", flags=[], created_by=ctx.account.id, updated_by=ctx.account.id,
    )
    ctx.db.add(it)
    ctx.db.flush()
    record(ctx.db, ctx.account.id, "create_kpi_item", "kpi_item", it.id, f"{emp.full_name}: {it.title}", new=body.model_dump())
    ctx.db.commit()
    return {"id": it.id}


class ItemPatch(BaseModel):
    version: int
    title: str | None = Field(default=None, min_length=2, max_length=500)
    kpi_type_id: int | None = None
    target_value: float | None = Field(default=None, ge=0)
    achieved_value: float | None = Field(default=None, ge=0)
    target_date: date | None = None
    outcome: str | None = Field(default=None, pattern="^(Completed|Not completed)$")
    completed_date: date | None = None
    delay_reason: str | None = None
    remarks: str | None = None


ITEM_FIELDS = ("title", "target_value", "achieved_value", "target_date", "outcome", "completed_date", "delay_reason", "remarks")


@router.patch("/items/{item_id}")
def update_item(item_id: int, body: ItemPatch, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    it = ctx.db.get(KpiItem, item_id)
    if it is None:
        raise NotFound("KPI row")
    month = ctx.db.get(KpiMonth, it.month_id)
    _require_open(month, it.employee_id, ctx.db)
    emp = ind.require_person(pilot, ctx.scope, it.employee_id)
    ind.require_edit(pilot, ctx.account, emp, month)
    if body.version != it.version:
        raise Conflict("Someone saved this row after you opened it. Compare before saving again.", current={"version": it.version, "achieved_value": it.achieved_value, "outcome": it.outcome})
    changes = body.model_dump(exclude_unset=True, exclude={"version"})
    if it.kind == "duty" and ({"target_value", "achieved_value"} & changes.keys()):
        raise RuleViolation("Duty tasks have no numeric target or achieved value.", "achieved_value")
    if "target_date" in changes:
        changes["target_date"] = _manager_due(changes["target_date"], month)
    if "kpi_type_id" in changes:
        if it.kind != "kpi" or not changes["kpi_type_id"] or not ctx.db.get(KpiType, changes["kpi_type_id"]):
            raise RuleViolation("Choose a valid KPI type.", "kpi_type_id")
    if "achieved_value" in changes and changes["achieved_value"] is None and not it.is_team:
        changes["achieved_value"] = 0.0  # a cleared achieved value goes back to the default of 0
    new_outcome = changes.get("outcome", it.outcome)
    new_completed = changes.get("completed_date", it.completed_date)
    if "outcome" in changes and new_outcome != "Completed":
        new_completed = None
        changes["completed_date"] = None
    if new_outcome == "Completed" and new_completed is None:
        raise RuleViolation("Enter the completed date.", "completed_date")
    if new_completed and new_completed > today():
        raise RuleViolation("The completed date cannot be in the future.", "completed_date")
    if new_completed and "outcome" not in changes and it.outcome is None:
        changes["outcome"] = "Completed"
    old = {k: getattr(it, k) for k in changes}
    for k, v in changes.items():
        setattr(it, k, v.strip() if isinstance(v, str) else v)
    if {"outcome", "completed_date", "delay_reason", "target_date"} & changes.keys() and it.delay_validated:
        it.delay_validated, it.delay_validated_by = False, None  # a changed delay needs validating again
    if "delay_reason" in changes:
        it.delay_reason_by = ctx.account.id if it.delay_reason else None
    it.updated_by = ctx.account.id
    it.version += 1
    record(ctx.db, ctx.account.id, "update_kpi_item", "kpi_item", it.id, f"{emp.full_name}: {it.title[:80]}", old=old, new=changes)
    ctx.db.commit()
    return {"id": it.id, "version": it.version}


def _can_validate_delay(ctx: Ctx, it: KpiItem, month: KpiMonth) -> bool:
    return (
        month.state == "Open"
        and ctx.account.role in ("hod", "admin")
        and bool(it.delay_reason)
        and not it.delay_validated
        and it.delay_reason_by != ctx.account.id
        and it.employee_id != ctx.account.employee_id
    )


class DecisionIn(BaseModel):
    decision: str
    reason: str = ""


@router.post("/items/{item_id}/validate-delay")
def validate_delay(item_id: int, body: DecisionIn, ctx: Ctx = Depends(get_ctx)):
    """HOD or Admin, never whoever entered the reason. A validated delay leaves On-time and the -2 penalty (D5)."""
    pilot = _pilot(ctx)
    it = ctx.db.get(KpiItem, item_id)
    if it is None:
        raise NotFound("KPI row")
    month = ctx.db.get(KpiMonth, it.month_id)
    ind.require_person(pilot, ctx.scope, it.employee_id)
    if not it.delay_reason:
        raise RuleViolation("There is no delay reason to validate.", "delay_reason")
    if not _can_validate_delay(ctx, it, month):
        raise Forbidden("Delay reasons are validated by the HOD or Admin, and never by whoever entered them.")
    if body.decision not in ("validate", "reject"):
        raise RuleViolation("Choose validate or reject.", "decision")
    if body.decision == "reject" and not body.reason.strip():
        raise RuleViolation("Give a reason for rejecting the delay reason.", "reason")
    it.delay_validated = body.decision == "validate"
    it.delay_validated_by = ctx.account.id if it.delay_validated else None
    if body.decision == "reject":
        it.flags = [*(it.flags or []), {"level": "warn", "text": f"Delay reason rejected: {body.reason.strip()}"}]
    it.version += 1
    record(ctx.db, ctx.account.id, f"delay_{body.decision}", "kpi_item", it.id, it.title[:80], new={"validated": it.delay_validated}, reason=body.reason or None)
    ctx.db.commit()
    return {"ok": True}


class ReasonIn(BaseModel):
    reason: str = Field(min_length=3)


@router.post("/items/{item_id}/remove")
def remove_item(item_id: int, body: ReasonIn, ctx: Ctx = Depends(get_ctx)):
    """Remove a KPI or duty task from this month only. Earlier months keep it; it stops carrying over. Kept in history."""
    pilot = _pilot(ctx)
    it = ctx.db.get(KpiItem, item_id)
    if it is None or it.removed_at is not None:
        raise NotFound("KPI row")
    month = ctx.db.get(KpiMonth, it.month_id)
    _require_open(month, it.employee_id, ctx.db)
    emp = ind.require_person(pilot, ctx.scope, it.employee_id)
    ind.require_edit(pilot, ctx.account, emp, month)
    it.removed_at, it.removed_by, it.removed_reason = now_ist(), ctx.account.id, body.reason.strip()
    it.version += 1
    record(ctx.db, ctx.account.id, "remove_kpi_item", "kpi_item", it.id, f"{emp.full_name}: {it.title[:80]}", reason=body.reason.strip())
    ctx.db.commit()
    return {"ok": True}


class FlagIn(BaseModel):
    index: int = Field(ge=0)
    note: str = Field(min_length=2)


@router.post("/items/{item_id}/resolve-warning")
def resolve_warning(item_id: int, body: FlagIn, ctx: Ctx = Depends(get_ctx)):
    """Mark an import or data warning as checked. Name-match warnings (D7) are confirmed by an Admin."""
    pilot = _pilot(ctx)
    it = ctx.db.get(KpiItem, item_id)
    if it is None:
        raise NotFound("KPI row")
    month = ctx.db.get(KpiMonth, it.month_id)
    _require_open(month, it.employee_id, ctx.db)
    emp = ind.require_person(pilot, ctx.scope, it.employee_id)
    flags = list(it.flags or [])
    if body.index >= len(flags) or flags[body.index].get("level") != "warn":
        raise RuleViolation("That warning does not exist.", "index")
    f = dict(flags[body.index])
    if f.get("resolved"):
        raise RuleViolation("This warning is already checked.", "index")
    name_match = f.get("text", "").startswith("Name match")
    if name_match and ctx.account.role != "admin":
        raise Forbidden("Name matches are confirmed by an Admin.")
    if not name_match and ctx.account.role != "admin":
        ind.require_edit(pilot, ctx.account, emp, month)
    f.update({"resolved": True, "note": body.note.strip(), "resolved_by": ctx.account.employee.full_name, "resolved_at": now_ist().isoformat(timespec="minutes")})
    flags[body.index] = f
    it.flags = flags
    record(ctx.db, ctx.account.id, "resolve_warning", "kpi_item", it.id, f"{emp.full_name}: {it.title[:60]}", new={"warning": f["text"]}, reason=body.note.strip())
    ctx.db.commit()
    return {"ok": True}


# --- training (TWCS) -----------------------------------------------------------------------------


@router.get("/months/{month_id}/training")
def month_training(month_id: int, employee_id: int | None = None, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    ids = ind.visible_ids(pilot, ctx.scope)
    if employee_id is not None:
        ind.require_person(pilot, ctx.scope, employee_id)
        ids = {employee_id}
    data = ind.load_month(pilot, month)
    cfg = data.rule.config
    end = month_end(month.month)
    by_emp = pilot.by_id()
    people, team_rows = [], []
    for emp_id in sorted(ids, key=lambda i: by_emp[i].code):
        emp = by_emp[emp_id]
        card = ind.compute(pilot, data, emp)
        t = card["training"]
        team_rows.extend(m for m in t["modules"])
        ok, _ = ind.can_edit(pilot, ctx.account, emp, month, data.evaluations.get(emp_id))
        mods = data.modules.get(emp_id, [])
        people.append({
            "employee": person(emp),
            "twcs": t["score"], "status": t["status"], "band": t["band"], "keka": t["keka"], "due": t["due"], "counted": t["counted"], "future": t["future"],
            "exempt": t["exempt"], "provisional": t["provisional"], "reason": t["reason"],
            "can_edit": ok,
            "modules": [{**ind.module_dict(m, month.training_cutoff, cfg), "in_month": month.month <= m.due_date <= end,
                         "can_verify": _can_verify(ctx, m), "can_decide_exemption": _can_decide_exemption(ctx, m)} for m in mods],
        })
    return {"month": _month_dict(ctx, month), "cutoff": month.training_cutoff, "tracks": list(ss.MODULE_TRACKS), "bands": cfg["twcs"]["bands"],
            "credits": cfg["twcs"]["credits"], "credit_beyond": cfg["twcs"]["credit_beyond"], "team_twcs": ss.team_twcs(team_rows), "people": people}


class ModuleIn(BaseModel):
    employee_id: int
    title: str = Field(min_length=2, max_length=200)
    track: str
    assigned_date: date
    due_date: date
    completed_date: date | None = None
    note: str | None = None


def _module_month(ctx: Ctx, due: date) -> KpiMonth | None:
    return ctx.db.scalar(select(KpiMonth).where(KpiMonth.month == due.replace(day=1)))


def _require_module_edit(ctx: Ctx, pilot: ind.Pilot, emp_id: int, due: date) -> None:
    emp = ind.require_person(pilot, ctx.scope, emp_id)
    km = _module_month(ctx, due)
    if km is not None:
        ind.require_edit(pilot, ctx.account, emp, km)
    elif not (ctx.account.role == "admin" or ((ev := pilot.evaluator(emp)) is not None and ev.id == ctx.account.id)):
        raise Forbidden("Only the evaluator or an Admin records this person's training.")


def _check_module_dates(assigned: date, due: date, completed: date | None) -> None:
    if due < assigned:
        raise RuleViolation("The due date cannot be before the assigned date.", "due_date")
    if completed and completed > today():
        raise RuleViolation("The completion date cannot be in the future.", "completed_date")
    if completed and completed < assigned:
        raise RuleViolation("The completion date cannot be before the assigned date.", "completed_date")


@router.post("/training")
def create_module(body: ModuleIn, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    if body.track not in ss.MODULE_TRACKS:
        raise RuleViolation("Choose a module track.", "track")
    _check_module_dates(body.assigned_date, body.due_date, body.completed_date)
    _require_module_edit(ctx, pilot, body.employee_id, body.due_date)
    m = TrainingModule(**body.model_dump(), status="Completed" if body.completed_date else "Assigned", created_by=ctx.account.id,
                       recorded_by=ctx.account.id if body.completed_date else None)
    ctx.db.add(m)
    ctx.db.flush()
    record(ctx.db, ctx.account.id, "create_training_module", "training_module", m.id, f"{m.title} ({m.track})", new=body.model_dump())
    ctx.db.commit()
    return {"id": m.id}


class ModulePatch(BaseModel):
    title: str | None = Field(default=None, min_length=2, max_length=200)
    track: str | None = None
    due_date: date | None = None
    completed_date: date | None = None
    attempts: int | None = Field(default=None, ge=1, le=20)
    note: str | None = None


@router.patch("/training/{module_id}")
def update_module(module_id: int, body: ModulePatch, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    m = ctx.db.get(TrainingModule, module_id)
    if m is None:
        raise NotFound("Training module")
    changes = body.model_dump(exclude_unset=True)
    if "track" in changes and changes["track"] not in ss.MODULE_TRACKS:
        raise RuleViolation("Choose a module track.", "track")
    _require_module_edit(ctx, pilot, m.employee_id, m.due_date)
    if "due_date" in changes:
        _require_module_edit(ctx, pilot, m.employee_id, changes["due_date"])
    if m.status == "Exempt":
        raise RuleViolation("This module has an approved exemption.", "status")
    _check_module_dates(m.assigned_date, changes.get("due_date", m.due_date), changes.get("completed_date", m.completed_date))
    old = {k: getattr(m, k) for k in changes}
    for k, v in changes.items():
        setattr(m, k, v)
    if "completed_date" in changes:
        m.status = "Completed" if m.completed_date else "Assigned"
        m.recorded_by = ctx.account.id if m.completed_date else None
        m.verified_by, m.verified_at = None, None  # a changed completion needs verifying again
    record(ctx.db, ctx.account.id, "update_training_module", "training_module", m.id, m.title, old=old, new=changes)
    ctx.db.commit()
    return {"id": m.id}


def _can_verify(ctx: Ctx, m: TrainingModule) -> bool:
    return (ctx.account.role in ("hod", "admin") and m.completed_date is not None and m.verified_by is None
            and m.recorded_by != ctx.account.id and m.employee_id != ctx.account.employee_id)


def _can_decide_exemption(ctx: Ctx, m: TrainingModule) -> bool:
    return (ctx.account.role in ("hod", "admin") and m.status == "Exemption requested"
            and m.exemption_requested_by != ctx.account.id and m.employee_id != ctx.account.employee_id)


@router.post("/training/{module_id}/verify")
def verify_module(module_id: int, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    m = ctx.db.get(TrainingModule, module_id)
    if m is None:
        raise NotFound("Training module")
    ind.require_person(pilot, ctx.scope, m.employee_id)
    if not _can_verify(ctx, m):
        raise Forbidden("Completions are verified by the HOD or Admin, never by whoever recorded them.")
    km = _module_month(ctx, m.due_date)
    if km is not None:
        _require_open(km, m.employee_id, ctx.db)
    m.verified_by, m.verified_at = ctx.account.id, now_ist()
    record(ctx.db, ctx.account.id, "verify_training_module", "training_module", m.id, m.title, new={"completed_date": m.completed_date})
    ctx.db.commit()
    return {"ok": True}


# --- training catalogue and team assignment -------------------------------------------------------


def _catalog_upsert(ctx: Ctx, title: str, track: str) -> TrainingCatalogItem:
    item = ctx.db.scalar(select(TrainingCatalogItem).where(TrainingCatalogItem.title == title))
    if item is None:
        item = TrainingCatalogItem(title=title, track=track, created_by=ctx.account.id)
        ctx.db.add(item)
        ctx.db.flush()
    elif not item.active:
        item.active = True
    return item


@router.get("/training-catalog")
def training_catalog(ctx: Ctx = Depends(get_ctx)):
    _pilot(ctx)
    return [{"id": c.id, "title": c.title, "track": c.track, "illustrative": c.illustrative}
            for c in ctx.db.scalars(select(TrainingCatalogItem).where(TrainingCatalogItem.active.is_(True)).order_by(TrainingCatalogItem.track, TrainingCatalogItem.title))]


class CatalogIn(BaseModel):
    title: str = Field(min_length=2, max_length=200)
    track: str


@router.post("/training-catalog")
def add_catalog_item(body: CatalogIn, ctx: Ctx = Depends(get_ctx)):
    ctx.scope.require_role("admin", "hod", "manager")
    if body.track not in ss.MODULE_TRACKS:
        raise RuleViolation("Choose a module track.", "track")
    item = _catalog_upsert(ctx, body.title.strip(), body.track)
    record(ctx.db, ctx.account.id, "add_training_catalog", "training_catalog", item.id, item.title, new={"track": item.track})
    ctx.db.commit()
    return {"id": item.id, "title": item.title, "track": item.track}


class AssignIn(BaseModel):
    title: str = Field(min_length=2, max_length=200)
    track: str
    assigned_date: date
    due_date: date
    employee_ids: list[int] = Field(min_length=1, max_length=200)


@router.post("/training/assign")
def assign_training(body: AssignIn, ctx: Ctx = Depends(get_ctx)):
    """Assign one module to several people in one step. Each person reports its own result; completion is recorded per person later."""
    pilot = _pilot(ctx)
    if body.track not in ss.MODULE_TRACKS:
        raise RuleViolation("Choose a module track.", "track")
    _check_module_dates(body.assigned_date, body.due_date, None)
    title = body.title.strip()
    by_emp = pilot.by_id()
    results = []
    for emp_id in dict.fromkeys(body.employee_ids):
        emp = by_emp.get(emp_id)
        try:
            if emp is None:
                raise NotFound("Person")
            _require_module_edit(ctx, pilot, emp_id, body.due_date)
            if ctx.db.scalar(select(TrainingModule.id).where(TrainingModule.employee_id == emp_id, TrainingModule.title == title, TrainingModule.due_date == body.due_date)):
                raise RuleViolation("Already assigned with this due date.", "title")
            m = TrainingModule(employee_id=emp_id, title=title, track=body.track, assigned_date=body.assigned_date, due_date=body.due_date, status="Assigned", created_by=ctx.account.id)
            ctx.db.add(m)
            ctx.db.flush()
            record(ctx.db, ctx.account.id, "create_training_module", "training_module", m.id, f"{m.title} ({m.track})", new={"employee": emp.full_name, "due_date": body.due_date})
            results.append({"employee_id": emp_id, "name": emp.full_name, "ok": True})
        except HTTPException as e:
            detail = e.detail if isinstance(e.detail, dict) else {"message": str(e.detail)}
            results.append({"employee_id": emp_id, "name": emp.full_name if emp else None, "ok": False, "message": detail.get("message")})
    if any(r["ok"] for r in results):  # the catalogue only learns modules that were actually assigned
        _catalog_upsert(ctx, title, body.track)
    ctx.db.commit()
    return {"done": sum(r["ok"] for r in results), "failed": sum(not r["ok"] for r in results), "results": results}


class ExemptionIn(BaseModel):
    reason: str = Field(min_length=5)


@router.post("/training/{module_id}/exemption")
def request_exemption(module_id: int, body: ExemptionIn, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    m = ctx.db.get(TrainingModule, module_id)
    if m is None:
        raise NotFound("Training module")
    _require_module_edit(ctx, pilot, m.employee_id, m.due_date)
    if m.status in ("Exempt", "Exemption requested"):
        raise RuleViolation(f"This module is already {m.status.lower()}.", "status")
    m.status, m.exemption_reason, m.exemption_requested_by = "Exemption requested", body.reason.strip(), ctx.account.id
    record(ctx.db, ctx.account.id, "request_training_exemption", "training_module", m.id, m.title, reason=body.reason)
    ctx.db.commit()
    return {"ok": True}


@router.post("/training/{module_id}/exemption-decision")
def decide_exemption(module_id: int, body: DecisionIn, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    m = ctx.db.get(TrainingModule, module_id)
    if m is None:
        raise NotFound("Training module")
    ind.require_person(pilot, ctx.scope, m.employee_id)
    if not _can_decide_exemption(ctx, m):
        raise Forbidden("Exemptions are decided by the HOD or Admin, never by whoever requested them.")
    if body.decision not in ("approve", "reject"):
        raise RuleViolation("Choose approve or reject.", "decision")
    if body.decision == "reject" and not body.reason.strip():
        raise RuleViolation("Give a reason for rejecting the exemption.", "reason")
    m.status = "Exempt" if body.decision == "approve" else ("Completed" if m.completed_date else "Assigned")
    m.exemption_decided_by = ctx.account.id
    record(ctx.db, ctx.account.id, f"exemption_{body.decision}", "training_module", m.id, m.title, reason=body.reason or None)
    ctx.db.commit()
    return {"ok": True}


# --- evaluation ----------------------------------------------------------------------------------


def _can_approve(ctx: Ctx, pilot: ind.Pilot, emp, ev: MonthlyEvaluation) -> tuple[bool, str]:
    if ctx.account.employee_id == emp.id:
        return False, "No one approves their own evaluation."
    if ev.submitted_by == ctx.account.id or ev.updated_by == ctx.account.id:
        return False, "The approver must be someone other than the evaluator who wrote it."
    roles = pilot.approver_roles(emp)
    if ctx.account.role not in roles:
        return False, "This evaluation is approved by " + (" or ".join(r.upper() if r == "hod" else r.title() for r in roles)) + "."
    return True, ""


def _evaluation_payload(ctx: Ctx, pilot: ind.Pilot, emp, month: KpiMonth) -> dict:
    data = ind.load_month(pilot, month)
    card = ind.compute(pilot, data, emp)
    ev = data.evaluations.get(emp.id)
    can_edit, edit_reason = ind.can_edit(pilot, ctx.account, emp, month, ev)
    approve_ok, approve_reason = _can_approve(ctx, pilot, emp, ev) if ev and ev.state == "Submitted" else (False, "")
    weights = data.rule.config["weights"].get(pilot.dept.category, {})
    types = _types(ctx)
    method = ind.method_visible(ctx.account)
    out = {
        **(card if method else ind.redact_card(card)),
        "month": _month_dict(ctx, month),
        "evaluation": ind.evaluation_dict(ev),
        # Only ratings weighted in the category are asked for (and required before submitting).
        "rated_metrics": [{"metric": m, "label": ss.SALES_METRIC_LABELS[m], "required": bool(weights.get(m)), **({"weight": weights.get(m, 0)} if method else {})} for m in ss.RATED_METRICS],
        "method_visible": method,
        "item_details": [{**ind.item_dict(it, data.rule.config, data.sums, types), "can_validate_delay": _can_validate_delay(ctx, it, month)} for it in data.items.get(emp.id, [])],
        "approver_roles": list(pilot.approver_roles(emp)),
        "submitted_by": account_person(ctx.db, ev.submitted_by) if ev else None,
        "decided_by": account_person(ctx.db, ev.decided_by) if ev else None,
        "can_edit": can_edit,
        "edit_reason": edit_reason,
        "can_submit": can_edit and ev is not None and ev.state in ind.EDITABLE_EVAL_STATES,
        "can_approve": approve_ok,
        "approve_reason": approve_reason,
        "corrections": [_correction_dict(ctx, c) for c in ctx.db.scalars(select(CorrectionRequest).where(
            CorrectionRequest.month_id == month.id, CorrectionRequest.employee_id == emp.id).order_by(CorrectionRequest.id.desc()))],
        "can_request_correction": month.state == "Locked" and ctx.account.role in ("hod", "admin") and not ctx.db.scalar(
            select(CorrectionRequest.id).where(CorrectionRequest.month_id == month.id, CorrectionRequest.employee_id == emp.id, CorrectionRequest.state.in_(("Requested", "Open")))),
    }
    if method:
        out.update({"weights": weights, "tiers": data.rule.config["tiers"], "tiers_proposed": data.rule.config.get("tiers_proposed", True), "penalties": data.rule.config["penalties"]})
    return out


@router.get("/evaluations/{employee_id}/{month_id}")
def get_evaluation(employee_id: int, month_id: int, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    emp = ind.require_person(pilot, ctx.scope, employee_id)
    month = ind.require_month(ctx.db, month_id)
    return _evaluation_payload(ctx, pilot, emp, month)


class EvaluationIn(BaseModel):
    version: int | None = None
    csat_rating: int | None = Field(default=None, ge=1, le=5)
    quality_rating: int | None = Field(default=None, ge=1, le=5)
    compliance_rating: int | None = Field(default=None, ge=1, le=5)
    collaboration_rating: int | None = Field(default=None, ge=1, le=5)
    innovation_rating: int | None = Field(default=None, ge=1, le=5)
    twcs_score: float | None = Field(default=None, ge=0, le=100)
    training_none_due: bool = False
    comments: str | None = None


@router.put("/evaluations/{employee_id}/{month_id}")
def save_evaluation(employee_id: int, month_id: int, body: EvaluationIn, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    emp = ind.require_person(pilot, ctx.scope, employee_id)
    month = ind.require_month(ctx.db, month_id)
    _require_open(month, emp.id, ctx.db)
    ev = ind.require_edit(pilot, ctx.account, emp, month)
    if ev is None:
        ev = MonthlyEvaluation(month_id=month.id, employee_id=emp.id, state="Draft", version=0)
        ctx.db.add(ev)
    elif body.version is not None and body.version != ev.version:
        raise Conflict("This evaluation changed after you opened it. Compare before saving again.", current=ind.evaluation_dict(ev))
    fields = body.model_dump(exclude={"version"}, exclude_unset=True)  # the rating sheet saves only the cells it changed
    old = {k: getattr(ev, k) for k in fields}
    for k, v in fields.items():
        setattr(ev, k, v.strip() if isinstance(v, str) else v)
    ev.updated_by = ctx.account.id
    ev.version = (ev.version or 0) + 1
    ctx.db.flush()
    record(ctx.db, ctx.account.id, "save_evaluation", "monthly_evaluation", ev.id, f"{emp.full_name} {ind.month_label(month.month)}", old=old, new=fields)
    ctx.db.commit()
    return _evaluation_payload(ctx, pilot, emp, month)


@router.post("/evaluations/{employee_id}/{month_id}/submit")
def submit_evaluation(employee_id: int, month_id: int, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    emp = ind.require_person(pilot, ctx.scope, employee_id)
    month = ind.require_month(ctx.db, month_id)
    _require_open(month, emp.id, ctx.db)
    ev = ind.require_edit(pilot, ctx.account, emp, month)
    if ev is None:
        raise RuleViolation("Save the evaluation ratings first.", "quality_rating")
    weights = ctx.db.get(ScoringRule, month.rule_id).config["weights"].get(pilot.dept.category, {})
    for metric in ss.RATED_METRICS:
        if weights.get(metric) and getattr(ev, f"{metric}_rating") is None:
            raise RuleViolation(f"Enter the {ss.SALES_METRIC_LABELS[metric]} rating before submitting.", f"{metric}_rating")
    ev.state, ev.submitted_by, ev.submitted_at = "Submitted", ctx.account.id, now_ist()
    ev.decided_by = ev.decided_at = ev.decision_note = None
    ev.version += 1
    label = f"{emp.full_name} {ind.month_label(month.month)}"
    record(ctx.db, ctx.account.id, "submit_evaluation", "monthly_evaluation", ev.id, label)
    for acc in ctx.db.scalars(select(Account).where(Account.active.is_(True), Account.role.in_(pilot.approver_roles(emp)))):
        if acc.id != ctx.account.id and (acc.role == "admin" or ctx.scope.perf_people is None or acc.employee_id in {h.id for h in pilot.heads}):
            notify(ctx.db, acc.id, f"Evaluation to approve: {label}", None, f"/individual/team/{emp.id}?month={month.id}")
    ctx.db.commit()
    return _evaluation_payload(ctx, pilot, emp, month)


@router.post("/evaluations/{employee_id}/{month_id}/decide")
def decide_evaluation(employee_id: int, month_id: int, body: DecisionIn, ctx: Ctx = Depends(get_ctx)):
    """Approve (also verifies the month's recorded training completions, D8) or return with a reason."""
    pilot = _pilot(ctx)
    emp = ind.require_person(pilot, ctx.scope, employee_id)
    month = ind.require_month(ctx.db, month_id)
    _decide(ctx, pilot, emp, month, body.decision, body.reason)
    ctx.db.commit()
    return _evaluation_payload(ctx, pilot, emp, month)


def _decide(ctx: Ctx, pilot: ind.Pilot, emp, month: KpiMonth, decision: str, reason: str) -> None:
    """One approval decision with every rule applied (shared by the single and the bulk endpoint). Raises on any rule."""
    _require_open(month, emp.id, ctx.db)
    ev = ind.get_evaluation(ctx.db, month, emp.id)
    if ev is None or ev.state != "Submitted":
        raise RuleViolation("Only a submitted check-in can be approved or returned.", "state")
    ok, why = _can_approve(ctx, pilot, emp, ev)
    if not ok:
        raise Forbidden(why)
    if decision not in ("approve", "return"):
        raise RuleViolation("Choose approve or return.", "decision")
    reason = (reason or "").strip()
    if decision == "return" and not reason:
        raise RuleViolation("Give the reason for returning it.", "reason")
    verified = 0
    if decision == "approve":
        start, end = month.month, month_end(month.month)
        for m in ctx.db.scalars(select(TrainingModule).where(TrainingModule.employee_id == emp.id, TrainingModule.due_date >= start, TrainingModule.due_date <= end)):
            if m.completed_date and m.verified_by is None and m.recorded_by != ctx.account.id:
                m.verified_by, m.verified_at = ctx.account.id, now_ist()
                verified += 1
    ev.state = "Approved" if decision == "approve" else "Returned"
    ev.decided_by, ev.decided_at, ev.decision_note = ctx.account.id, now_ist(), reason or None
    ev.version += 1
    label = f"{emp.full_name} {ind.month_label(month.month)}"
    corr = ind.open_correction(ctx.db, month, emp.id)
    if corr is not None and decision == "approve":
        ctx.db.flush()
        ind.refreeze_person(pilot, month, emp, corr)
        record(ctx.db, ctx.account.id, "complete_correction", "correction", corr.id, label, old=corr.old_result, new=corr.new_result, reason=corr.reason)
    record(ctx.db, ctx.account.id, f"{decision}_evaluation", "monthly_evaluation", ev.id, label, new={"state": ev.state, "training_verified": verified}, reason=reason or None)
    notify(ctx.db, ev.submitted_by, f"Check-in {ev.state.lower()}: {label}", reason or None, f"/individual/team/{emp.id}?month={month.id}")


class DecideManyIn(BaseModel):
    employee_ids: list[int] = Field(min_length=1, max_length=200)
    decision: str
    reason: str = ""


@router.post("/months/{month_id}/decide-many")
def decide_many(month_id: int, body: DecideManyIn, ctx: Ctx = Depends(get_ctx)):
    """Approve or return several check-ins at once. Same rules as one by one; each person reports its own result."""
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    by_emp = pilot.by_id()
    results = []
    for emp_id in dict.fromkeys(body.employee_ids):
        emp = by_emp.get(emp_id)
        try:
            if emp is None:
                raise NotFound("Person")
            ind.require_person(pilot, ctx.scope, emp_id)
            _decide(ctx, pilot, emp, month, body.decision, body.reason)  # every rule is checked before anything changes
            results.append({"employee_id": emp_id, "name": emp.full_name, "ok": True})
        except HTTPException as e:
            detail = e.detail if isinstance(e.detail, dict) else {"message": str(e.detail)}
            results.append({"employee_id": emp_id, "name": emp.full_name if emp else None, "ok": False, "message": detail.get("message")})
    ctx.db.commit()
    return {"done": sum(r["ok"] for r in results), "failed": sum(not r["ok"] for r in results), "results": results}


# --- month lock ----------------------------------------------------------------------------------


@router.get("/months/{month_id}/lock-check")
def lock_check(month_id: int, ctx: Ctx = Depends(get_ctx)):
    ctx.scope.require_role("hod", "admin")
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    data = ind.load_month(pilot, month)
    not_approved, readiness = [], []
    for emp in pilot.people:
        ev = data.evaluations.get(emp.id)
        if ev is None or ev.state != "Approved":
            not_approved.append({"employee": person(emp), "state": ev.state if ev else "Not started"})
    pending_delays = sum(1 for rows in data.items.values() for it in rows if it.delay_reason and not it.delay_validated and ind.item_timing(it, data.rule.config) == "Delayed")
    pending_exemptions = sum(1 for mods in data.modules.values() for m in mods if m.status == "Exemption requested" and month.month <= m.due_date <= month_end(month.month))
    if pending_delays:
        readiness.append(f"{pending_delays} delay reason(s) not yet validated; they count as delays at lock.")
    if pending_exemptions:
        readiness.append(f"{pending_exemptions} training exemption request(s) undecided; those modules count as due.")
    if today() <= month.training_cutoff:
        readiness.append(f"The training cut-off ({month.training_cutoff.strftime('%d %b %Y')}) has not passed; completions until then would still count.")
    return {"month": _month_dict(ctx, month, pilot), "people": len(pilot.people), "approved": len(pilot.people) - len(not_approved),
            "not_approved": not_approved, "readiness": readiness, "rule": ind.rule_dict(data.rule) if ind.method_visible(ctx.account) else {"label": data.rule.label},
            "can_lock": month.state == "Open"}


class LockIn(BaseModel):
    note: str = Field(min_length=3)
    exclude_reason: str | None = None


@router.post("/months/{month_id}/lock")
def lock_month(month_id: int, body: LockIn, ctx: Ctx = Depends(get_ctx)):
    ctx.scope.require_role("hod", "admin")
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    _require_open(month)
    pending = [e for e in pilot.people if (ev := ind.get_evaluation(ctx.db, month, e.id)) is None or ev.state != "Approved"]
    if pending and not (body.exclude_reason or "").strip():
        raise RuleViolation(f"{len(pending)} evaluation(s) are not approved. Approve them, or give a reason to lock with them recorded as not evaluated.", "exclude_reason")
    counts = ind.snapshot_month(pilot, month, ctx.account, (body.exclude_reason or "").strip() or None)
    month.state, month.locked_by, month.locked_at, month.lock_note = "Locked", ctx.account.id, now_ist(), body.note.strip()
    record(ctx.db, ctx.account.id, "lock_month", "kpi_month", month.id, ind.month_label(month.month), new=counts, reason=body.note)
    ctx.db.commit()
    return {"month": _month_dict(ctx, month, pilot), "counts": counts}


# --- rewards -------------------------------------------------------------------------------------


@router.get("/rewards")
def rewards(month_id: int | None = None, ctx: Ctx = Depends(get_ctx)):
    """Monthly tiers from locked snapshots only. Open months show nothing: scores are not final until locked."""
    pilot = _pilot(ctx)
    locked = list(ctx.db.scalars(select(KpiMonth).where(KpiMonth.state == "Locked").order_by(KpiMonth.month.desc())))
    month = ctx.db.get(KpiMonth, month_id) if month_id else (locked[0] if locked else None)
    if month is None or month.state != "Locked":
        return {"months": [_month_dict(ctx, m) for m in locked], "month": _month_dict(ctx, month) if month else None, "rows": [], "tiers": None}
    ids = ind.visible_ids(pilot, ctx.scope)
    rule = ctx.db.get(ScoringRule, month.rule_id)
    method = ind.method_visible(ctx.account)
    rows = []
    for s in ctx.db.scalars(select(ScorecardSnapshot).where(ScorecardSnapshot.month_id == month.id)):
        if s.employee_id not in ids:
            continue
        emp = pilot.by_id().get(s.employee_id)
        gates = [
            {"gate": f"Coverage >= {rule.config['coverage_min']}%" if method else "Enough data to score", "result": "pass" if s.coverage >= rule.config["coverage_min"] else "fail", "source": "Locked scorecard"},
            {"gate": "Evaluation approved", "result": "pass" if s.status != "not_evaluated" else "fail", "source": "Monthly evaluation"},
            {"gate": "Training not in Support band", "result": "unknown" if s.twcs is None else ("pass" if s.twcs >= rule.config["penalties"]["training_threshold"] else "fail"), "source": "TWCS"},
        ]
        row = {"employee": person(emp), "final": s.final, "base": s.base, "coverage": s.coverage, "twcs": s.twcs, "band": s.data.get("training", {}).get("band"),
               "tier": s.tier, "status": s.status, "gates": gates, "eligible": s.tier is not None and all(g["result"] == "pass" for g in gates[:2])}
        if not method:
            row.pop("base")
            row.pop("coverage")
        rows.append(row)
    rows.sort(key=lambda r: (r["final"] is None, -(r["final"] or 0)))
    tiers = {name: sum(1 for r in rows if r["tier"] == name and r["eligible"]) for name, _ in rule.config["tiers"]}
    return {"months": [_month_dict(ctx, m) for m in locked], "month": _month_dict(ctx, month), "rows": rows, "tiers": tiers,
            "tier_rules": rule.config["tiers"] if method else None, "tiers_proposed": rule.config.get("tiers_proposed", True)}


class PublishIn(BaseModel):
    note: str = Field(min_length=3)


@router.post("/months/{month_id}/publish")
def publish(month_id: int, body: PublishIn, ctx: Ctx = Depends(get_ctx)):
    """Publishing is a separate, explicit action after lock. Published results show names and tiers only."""
    ctx.scope.require_role("hod", "admin")
    month = ind.require_month(ctx.db, month_id)
    if month.state != "Locked":
        raise RuleViolation("Lock the month before publishing recognition.", "state")
    if month.published_at:
        raise RuleViolation("This month's recognition is already published.", "state")
    month.published_by, month.published_at = ctx.account.id, now_ist()
    record(ctx.db, ctx.account.id, "publish_recognition", "kpi_month", month.id, ind.month_label(month.month), reason=body.note)
    ctx.db.commit()
    return _month_dict(ctx, month)


# --- home: quarterly performance dashboard -------------------------------------------------------

QUARTER_MONTHS = {1: (1, 3), 2: (4, 6), 3: (7, 9), 4: (10, 12)}
ATTENTION_KPI_PCT = 80  # KPI average below this is flagged
SHORT_MONTH = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


def _quarter_of(d: date) -> int:
    return (d.month - 1) // 3 + 1


def _quarter_label(year: int, q: int) -> str:
    a, b = QUARTER_MONTHS[q]
    return f"Q{q} {year} ({SHORT_MONTH[a - 1]}–{SHORT_MONTH[b - 1]})"


def _mean(values) -> float | None:
    v = [x for x in values if x is not None]
    return sum(v) / len(v) if v else None


def _group_stats(people: list[dict]) -> dict:
    """Averages over people who have a value; counts say how many contributed, so missing data is never read as zero."""
    finals = [p["quarter"]["final_avg"] for p in people]
    kpis = [p["quarter"]["kpi_avg"] for p in people]
    twcs = [p["quarter"]["twcs_avg"] for p in people]
    cells = [m for p in people for m in p["months"]]
    return {
        "people": len(people),
        "final_avg": _mean(finals), "final_n": sum(x is not None for x in finals),
        "kpi_avg": _mean(kpis), "kpi_n": sum(x is not None for x in kpis),
        "twcs_avg": _mean(twcs), "twcs_n": sum(x is not None for x in twcs),
        "checkins_total": len(cells),
        "checkins_approved": sum(m["evaluation_state"] == "Approved" for m in cells),
        "checkins_submitted": sum(m["evaluation_state"] == "Submitted" for m in cells),
        "attention": sum(bool(p["attention"]) for p in people),
    }


@router.get("/dashboard")
def dashboard(year: int | None = None, quarter: int | None = Query(default=None, ge=1, le=4), ctx: Ctx = Depends(get_ctx)):
    """Role-based quarterly dashboard for the Individual Tier Home.

    Quarter-to-date values come from the existing monthly scorecards (same scoring, same scope and redaction):
    a person's quarter score is the average of their scored months, and every average says how many people or months it covers.
    """
    pilot = _pilot(ctx)
    all_months = list(ctx.db.scalars(select(KpiMonth).order_by(KpiMonth.month)))
    if year is None or quarter is None:
        ref = today().replace(day=1)
        cands = [m for m in all_months if m.month <= ref] or all_months
        ref = cands[-1].month if cands else ref
        year, quarter = year or ref.year, quarter or _quarter_of(ref)
    a, b = QUARTER_MONTHS[quarter]
    q_months = [m for m in all_months if m.month.year == year and a <= m.month.month <= b]
    available = sorted({(m.month.year, _quarter_of(m.month)) for m in all_months}, reverse=True)

    # Per month, per person: the existing scorecard rows (already scoped to the viewer and redacted for non-Admins).
    per_month = [(m, {r["employee"]["id"]: r for r in _people_rows(ctx, pilot, m)}) for m in q_months]
    ids = ind.visible_ids(pilot, ctx.scope)
    by_emp = pilot.by_id()
    people = []
    for emp_id in sorted(ids, key=lambda i: by_emp[i].full_name):
        months, base = [], None
        for m, rows in per_month:
            r = rows.get(emp_id)
            if r is None:
                continue
            base = base or r
            months.append({
                "month_id": m.id, "label": ind.month_label(m.month), "state": m.state, "final": r["final"], "status": r["status"], "tier": r["tier"],
                "kpi_avg": r["achievement_pct"], "kpis_total": r["kpis_total"], "twcs": r["twcs"], "band": r["band"],
                "on_time": r["on_time"]["on_time"], "delayed": r["on_time"]["delayed"], "evaluation_state": r["evaluation_state"],
            })
        if base is None:
            base = {"employee": person(by_emp[emp_id]), "manager": person(by_emp[emp_id].manager) if by_emp[emp_id].manager else None, "evaluator": None}
        quarter_row = {
            "final_avg": _mean(x["final"] for x in months), "scored_months": sum(x["final"] is not None for x in months),
            "kpi_avg": _mean(x["kpi_avg"] for x in months), "twcs_avg": _mean(x["twcs"] for x in months),
            "delayed": sum(x["delayed"] for x in months), "months": len(months),
        }
        attention = []
        if quarter_row["kpi_avg"] is not None and quarter_row["kpi_avg"] < ATTENTION_KPI_PCT:
            attention.append(f"KPI average {quarter_row['kpi_avg']:.1f}% is below {ATTENTION_KPI_PCT}%")
        if any(x["band"] == "Support" for x in months):
            attention.append("Training in the Support band (TWCS below 50)")
        if quarter_row["delayed"]:
            attention.append(f"{quarter_row['delayed']} delayed task(s) this quarter")
        ended = [x for x in months if x["month_id"] in {m.id for m in q_months if m.month < today().replace(day=1)}]
        pending = [x["label"] for x in ended if x["evaluation_state"] in ("Not started", "Draft", "Returned")]
        if pending:
            attention.append(f"Check-in not submitted: {', '.join(pending)}")
        if months and quarter_row["scored_months"] == 0:
            attention.append("Not enough data to score any month yet")
        if not any(x["kpis_total"] for x in months):
            attention.append("No KPI targets set")
        people.append({"employee": base["employee"], "manager": base.get("manager"), "evaluator": base.get("evaluator"),
                       "months": months, "quarter": quarter_row, "attention": attention})

    monthly = []
    for m, rows in per_month:
        rs = [r for r in rows.values() if r["employee"]["id"] in ids]
        finals = [r["final"] for r in rs]
        monthly.append({
            "month_id": m.id, "label": ind.month_label(m.month), "state": m.state, "people": len(rs),
            "final_avg": _mean(finals), "final_n": sum(x is not None for x in finals),
            "kpi_avg": _mean(r["achievement_pct"] for r in rs), "twcs_avg": _mean(r["twcs"] for r in rs), "twcs_n": sum(r["twcs"] is not None for r in rs),
            "approved": sum(r["evaluation_state"] == "Approved" for r in rs), "submitted": sum(r["evaluation_state"] == "Submitted" for r in rs),
        })

    role = ctx.account.role
    out = {
        "role": role, "year": year, "quarter": quarter, "label": _quarter_label(year, quarter),
        "quarters": [{"year": y, "quarter": q, "label": _quarter_label(y, q)} for y, q in available],
        "months": [{"id": m.id, "label": ind.month_label(m.month), "state": m.state} for m in q_months],
        "summary": _group_stats(people),
        "monthly": monthly,
        "people": people,
        "awaiting_my_approval": sum(1 for _, rows in per_month for r in rows.values() if r["can_approve"]),
        "scope_label": "Your reporting team" if role == "manager" else (f"{pilot.dept.name} department" if role == "hod" else "Organisation (Individual Tier pilot)"),
    }
    if role in ("hod", "admin"):
        # Manager-wise comparison: people grouped by who checks them in.
        groups: dict[int, list[dict]] = {}
        names: dict[int, dict] = {}
        for p in people:
            ev = p["evaluator"]
            if ev:
                groups.setdefault(ev["id"], []).append(p)
                names[ev["id"]] = ev
        out["managers"] = sorted(
            [{"manager": names[k], "is_me": k == ctx.account.employee_id, **_group_stats(v)} for k, v in groups.items()],
            key=lambda g: g["manager"]["name"],
        )
    if role == "admin":
        out["departments"] = [{"department": pilot.dept.name, "category": pilot.dept.category, "in_pilot": True, **_group_stats(people)}]
        out["other_departments"] = len(list(ctx.db.scalars(select(Department.id)))) - 1
        trend = []
        for q in range(1, 5):
            qa, qb = QUARTER_MONTHS[q]
            ms = [m for m in all_months if m.month.year == year and qa <= m.month.month <= qb]
            rs = [r for m in ms for r in _people_rows(ctx, pilot, m)] if q != quarter else [r for _, rows in per_month for r in rows.values()]
            finals = [r["final"] for r in rs]
            trend.append({"quarter": q, "label": f"Q{q}", "months": len(ms), "final_avg": _mean(finals), "final_n": sum(x is not None for x in finals),
                          "kpi_avg": _mean(r["achievement_pct"] for r in rs), "twcs_avg": _mean(r["twcs"] for r in rs)})
        out["quarterly_trend"] = trend
    return out


# --- reports and exports -------------------------------------------------------------------------


@router.get("/reports")
def reports(month_id: int, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    rows = _people_rows(ctx, pilot, month)
    data = ind.load_month(pilot, month)
    ids = {r["employee"]["id"] for r in rows}
    types = _types(ctx)
    by_type: dict[int, dict] = {}
    gaps = []
    for emp_id in ids:
        for it in data.items.get(emp_id, []):
            if it.kind == "kpi" and it.kpi_type_id and not it.is_team:
                t = by_type.setdefault(it.kpi_type_id, {"type": types[it.kpi_type_id].name, "unit": types[it.kpi_type_id].unit, "rows": 0, "target": 0.0, "achieved": 0.0, "with_actual": 0})
                t["rows"] += 1
                if it.target_value is not None and it.achieved_value is not None:
                    t["target"] += it.target_value
                    t["achieved"] += it.achieved_value
                    t["with_actual"] += 1
            for f in it.flags or []:
                if f.get("level") == "warn" and not f.get("resolved"):
                    gaps.append({"employee": person(pilot.by_id()[emp_id]), "row": it.title[:90], "source_row": it.source_row, "issue": f["text"]})
    for r in rows:
        if r["kpi_rows"] == 0:
            gaps.append({"employee": r["employee"], "row": None, "source_row": None, "issue": "No targets set for this month."})
    trend = []
    for m in ctx.db.scalars(select(KpiMonth).order_by(KpiMonth.month)):
        mrows = rows if m.id == month.id else _people_rows(ctx, pilot, m)
        ach = [r["achievement_pct"] for r in mrows if r["achievement_pct"] is not None]
        fin = [r["final"] for r in mrows if r["final"] is not None]
        tw = [r["twcs"] for r in mrows if r["twcs"] is not None]
        trend.append({"month": ind.month_label(m.month), "state": m.state, "achievement_avg": sum(ach) / len(ach) if ach else None, "achievement_n": len(ach),
                      "final_avg": sum(fin) / len(fin) if fin else None, "final_n": len(fin), "twcs_avg": sum(tw) / len(tw) if tw else None})
    team_mod_rows = [mr for emp_id in ids for mr in ind.compute(pilot, data, pilot.by_id()[emp_id])["training"]["modules"]]
    return {
        "month": _month_dict(ctx, month, pilot),
        "people": rows,
        "by_type": sorted(by_type.values(), key=lambda t: t["type"]),
        "on_time": {k: sum(r["on_time"][k] for r in rows) for k in ("on_time", "delayed", "delay_validated", "no_outcome")},
        "team_twcs": ss.team_twcs(team_mod_rows),
        "gaps": gaps,
        "trend": trend,
    }


def _csv(rows: list[list], filename: str) -> Response:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return Response(buf.getvalue(), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _fmt(v, digits=2):
    return "" if v is None else (f"{v:.{digits}f}" if isinstance(v, float) else v)


@router.get("/reports/export.csv")
def export_csv(month_id: int, ctx: Ctx = Depends(get_ctx)):
    """Scoped to the viewer; unavailable values stay blank with the reason column, never 0."""
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    rows = _people_rows(ctx, pilot, month)
    rule = ctx.db.get(ScoringRule, month.rule_id)
    method = ind.method_visible(ctx.account)
    out = [["Month", ind.month_label(month.month), "State", month.state, "Rule", rule.label, "Exported", now_ist().strftime("%d %b %Y %H:%M IST")],
           ["Employee ID", "Name", "Check-in by", "Check-in", "Achievement %", "On time", "Delayed", "TWCS", "Band", "KEKA rating"]
           + (["Coverage %", "Base", "Penalty"] if method else []) + ["Final", "Tier", "Status", "Readiness"]]
    for r in rows:
        ent = r["check_in"]["entered"]
        out.append([r["employee"]["code"], r["employee"]["name"], (r["evaluator"] or {}).get("name", ""), r["evaluation_state"],
                    _fmt(ent.get("achievement_pct") if ent.get("achievement_pct") is not None else r["achievement_pct"]),
                    ent.get("on_time_count") if ent.get("on_time_count") is not None else r["on_time"]["on_time"],
                    ent.get("delayed_count") if ent.get("delayed_count") is not None else r["on_time"]["delayed"],
                    _fmt(r["twcs"]), r["band"] or "", "" if r["keka"] is None else r["keka"]]
                   + ([_fmt(r["coverage"], 0), _fmt(r["base"]), r["penalty_total"]] if method else [])
                   + [_fmt(r["final"]), r["tier"] or "", "Insufficient data" if r["status"] == "insufficient" else r["status"], "; ".join(r["readiness"])])
    record(ctx.db, ctx.account.id, "export_report", "kpi_month", month.id, ind.month_label(month.month), new={"rows": len(rows)})
    ctx.db.commit()
    return _csv(out, f"finvasia-elevate-sales-{month.month:%Y-%m}.csv")


@router.get("/keka-export.csv")
def keka_export(month_id: int, ctx: Ctx = Depends(get_ctx)):
    """Monthly training ratings for upload to KEKA. There is no live KEKA connection."""
    ctx.scope.require_role("hod", "admin")
    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    rows = _people_rows(ctx, pilot, month)
    out = [["Employee ID", "HRMS code", "Name", "Month", "TWCS", "Band", "KEKA rating", "Modules counted", "Status"]]
    by_emp = pilot.by_id()
    for r in rows:
        emp = by_emp[r["employee"]["id"]]
        out.append([emp.code, emp.hrms_code or "", emp.full_name, ind.month_label(month.month), _fmt(r["twcs"]), r["band"] or "", "" if r["keka"] is None else r["keka"],
                    "", {"scored": "Scored", "na": "No training due", "missing": "Not recorded"}.get(r["twcs_status"], r["twcs_status"])])
    record(ctx.db, ctx.account.id, "export_keka", "kpi_month", month.id, ind.month_label(month.month), new={"rows": len(rows), "state": month.state})
    ctx.db.commit()
    return _csv(out, f"keka-training-{month.month:%Y-%m}.csv")


# --- KPI Scoring (scoring method) ----------------------------------------------------------------


@router.get("/rules")
def rules(ctx: Ctx = Depends(get_ctx)):
    ctx.scope.require_role("admin")
    rows = list(ctx.db.scalars(select(ScoringRule).order_by(ScoringRule.id.desc())))
    used = {}
    for m in ctx.db.scalars(select(KpiMonth)):
        used.setdefault(m.rule_id, []).append({"label": ind.month_label(m.month), "state": m.state})
    active = ctx.db.scalar(select(ScoringRule).where(ScoringRule.state == "Approved").order_by(ScoringRule.effective_month.desc(), ScoringRule.id.desc()))
    imports = [{"source": b.source, "at": b.created_at, "summary": b.summary} for b in ctx.db.scalars(select(ImportBatch).where(ImportBatch.source.like("%Sales%")).order_by(ImportBatch.id.desc()))]
    return {
        "active_id": active.id if active else None,
        "rules": [{**ind.rule_dict(r), "author": account_person(ctx.db, r.author_account_id), "approver": account_person(ctx.db, r.approver_account_id), "months": used.get(r.id, []),
                   "can_edit": ctx.account.role == "admin" and r.state in ("Draft", "Returned") and r.author_account_id == ctx.account.id,
                   "can_decide": r.state == "Pending approval" and r.author_account_id != ctx.account.id} for r in rows],
        "metrics": [{"metric": m, "label": ss.SALES_METRIC_LABELS[m]} for m in ss.SALES_METRICS],
        "kpi_types": kpi_types(ctx),
        "imports": imports,
        "can_author": ctx.account.role == "admin",
    }


class RuleIn(BaseModel):
    label: str = Field(min_length=2, max_length=40)
    summary: str | None = None
    effective_month: date
    config: dict


def _validate_config(cfg: dict) -> None:
    base = ss.DEFAULT_SALES_CONFIG
    for key in ("weights", "anchors", "penalties", "twcs", "tiers", "coverage_min", "achievement_cap", "rating_multiplier", "min_on_time_outcomes"):
        if key not in cfg:
            raise RuleViolation(f"Missing setting: {key}.", key)
    for cat, w in cfg["weights"].items():
        bad = set(w) - set(ss.SALES_METRICS)
        if bad:
            raise RuleViolation(f"{cat}: unknown metric {', '.join(sorted(bad))}.", "weights")
        total = sum(float(v or 0) for v in w.values())
        if abs(total - 100) > 0.001:
            raise RuleViolation(f"{cat} weights total {total:g}%. Each category must total 100%.", "weights")
        if any(float(v or 0) < 0 for v in w.values()):
            raise RuleViolation(f"{cat}: weights cannot be negative.", "weights")
    for name in ("achievement", "on_time"):
        pts = cfg["anchors"].get(name) or []
        if len(pts) < 2 or any(len(p) != 2 for p in pts):
            raise RuleViolation(f"The {name.replace('_', '-')} curve needs at least two [raw, score] points.", "anchors")
        if any(not 0 <= p[1] <= 100 for p in pts):
            raise RuleViolation("Curve scores must be between 0 and 100.", "anchors")
    if not 0 < float(cfg["coverage_min"]) <= 100:
        raise RuleViolation("Minimum coverage must be between 1 and 100.", "coverage_min")
    tiers = cfg["tiers"]
    if any(a[1] <= b[1] for a, b in zip(tiers, tiers[1:])):
        raise RuleViolation("Tier thresholds must decrease from the first tier to the last.", "tiers")
    bands = cfg["twcs"]["bands"]
    if any(a[0] <= b[0] for a, b in zip(bands, bands[1:])):
        raise RuleViolation("TWCS band minimums must decrease.", "twcs")
    for k in base["penalties"]:
        if k not in cfg["penalties"]:
            raise RuleViolation(f"Missing penalty setting: {k}.", "penalties")


@router.post("/rules")
def create_rule(body: RuleIn, ctx: Ctx = Depends(get_ctx)):
    ctx.scope.require_role("admin")
    _validate_config(body.config)
    r = ScoringRule(label=body.label.strip(), summary=body.summary, effective_month=body.effective_month.replace(day=1), config=body.config, state="Draft", author_account_id=ctx.account.id)
    ctx.db.add(r)
    ctx.db.flush()
    record(ctx.db, ctx.account.id, "create_scoring_rule", "scoring_rule", r.id, r.label, new={"effective_month": r.effective_month})
    ctx.db.commit()
    return ind.rule_dict(r)


@router.put("/rules/{rule_id}")
def update_rule(rule_id: int, body: RuleIn, ctx: Ctx = Depends(get_ctx)):
    ctx.scope.require_role("admin")
    r = ctx.db.get(ScoringRule, rule_id)
    if r is None:
        raise NotFound("Scoring rule")
    if r.state not in ("Draft", "Returned") or r.author_account_id != ctx.account.id:
        raise Forbidden("Only the author edits a draft or returned rule. Approved rules never change; make a new version.")
    _validate_config(body.config)
    old = {"config": r.config, "label": r.label}
    r.label, r.summary, r.effective_month, r.config, r.state = body.label.strip(), body.summary, body.effective_month.replace(day=1), copy.deepcopy(body.config), "Draft"
    record(ctx.db, ctx.account.id, "update_scoring_rule", "scoring_rule", r.id, r.label, old=old, new={"config": r.config})
    ctx.db.commit()
    return ind.rule_dict(r)


@router.get("/rules/{rule_id}/preview")
def preview_rule(rule_id: int, month_id: int, ctx: Ctx = Depends(get_ctx)):
    """Impact preview: every visible person's result under the month's own rule and under this one. Nothing is saved."""
    ctx.scope.require_role("admin")
    pilot = _pilot(ctx)
    r = ctx.db.get(ScoringRule, rule_id)
    if r is None:
        raise NotFound("Scoring rule")
    month = ind.require_month(ctx.db, month_id)
    data = ind.load_month(pilot, month)
    ids = ind.visible_ids(pilot, ctx.scope)
    out = []
    for emp in pilot.people:
        if emp.id not in ids:
            continue
        cur = ind.compute(pilot, data, emp, cfg=data.rule.config)["scorecard"]
        new = ind.compute(pilot, data, emp, cfg=r.config)["scorecard"]
        out.append({"employee": person(emp), "current": {"final": cur["final"], "tier": cur["tier"], "coverage": cur["coverage"], "status": cur["status"]},
                    "proposed": {"final": new["final"], "tier": new["tier"], "coverage": new["coverage"], "status": new["status"]},
                    "delta": None if cur["final"] is None or new["final"] is None else new["final"] - cur["final"]})
    changed = sum(1 for o in out if o["current"]["tier"] != o["proposed"]["tier"] or o["current"]["status"] != o["proposed"]["status"])
    return {"month": _month_dict(ctx, month), "current_rule": data.rule.label, "rule": r.label, "rows": out, "tier_changes": changed,
            "note": "Locked months keep their stored results; this preview never changes them." if month.state == "Locked" else ""}


@router.post("/rules/{rule_id}/submit")
def submit_rule(rule_id: int, ctx: Ctx = Depends(get_ctx)):
    ctx.scope.require_role("admin")
    r = ctx.db.get(ScoringRule, rule_id)
    if r is None or r.author_account_id != ctx.account.id or r.state not in ("Draft", "Returned"):
        raise Forbidden("Only the author submits a draft rule.")
    r.state = "Pending approval"
    record(ctx.db, ctx.account.id, "submit_scoring_rule", "scoring_rule", r.id, r.label)
    for acc in ctx.db.scalars(select(Account).where(Account.active.is_(True), Account.role == "admin", Account.id != ctx.account.id)):
        notify(ctx.db, acc.id, f"Scoring rule to approve: {r.label}", r.summary, "/individual/kpi-scoring")
    ctx.db.commit()
    return ind.rule_dict(r)


@router.post("/rules/{rule_id}/decide")
def decide_rule(rule_id: int, body: DecisionIn, ctx: Ctx = Depends(get_ctx)):
    """A second, independent Admin activates the rule. Open months from its effective month switch to it; locked months keep theirs.
    The scoring method is Admin-only, so the HOD does not approve it."""
    ctx.scope.require_role("admin")
    r = ctx.db.get(ScoringRule, rule_id)
    if r is None or r.state != "Pending approval":
        raise RuleViolation("Only a rule waiting for approval can be decided.", "state")
    if r.author_account_id == ctx.account.id:
        raise Forbidden("The author cannot approve their own rule.")
    if body.decision not in ("approve", "return"):
        raise RuleViolation("Choose approve or return.", "decision")
    if body.decision == "return" and not body.reason.strip():
        raise RuleViolation("Give the reason for returning it.", "reason")
    r.state = "Approved" if body.decision == "approve" else "Returned"
    r.approver_account_id, r.decided_at, r.decision_note = ctx.account.id, now_ist(), body.reason.strip() or None
    switched = []
    if r.state == "Approved":
        for m in ctx.db.scalars(select(KpiMonth).where(KpiMonth.state == "Open", KpiMonth.month >= r.effective_month)):
            m.rule_id = r.id
            m.training_cutoff = training_cutoff(m.month, r.config)
            switched.append(ind.month_label(m.month))
    record(ctx.db, ctx.account.id, f"{body.decision}_scoring_rule", "scoring_rule", r.id, r.label, new={"state": r.state, "months": switched}, reason=body.reason or None)
    notify(ctx.db, r.author_account_id, f"Scoring rule {r.state.lower()}: {r.label}", body.reason or None, "/individual/kpi-scoring")
    ctx.db.commit()
    return {**ind.rule_dict(r), "months_switched": switched}


# --- Excel template and upload for the monthly check-in (R&R parameter columns) -----------------

# (column header, evaluation field, kind). Headers follow the R&R Parameters sheet, plus Training. Tasks On-Time, Tasks Delayed and
# Achievement % are worked out from the KPIs, so they are not uploaded; only ratings weighted in the category are offered.
AUTO_FIELDS = {"on_time_count": "Tasks On-Time", "delayed_count": "Tasks Delayed", "achievement_pct": "Achievement %"}
SHEET_FIELDS = [
    ("Quality", "quality_rating", "rating"),
    ("Collaboration", "collaboration_rating", "rating"),
    ("Innovation", "innovation_rating", "rating"),
    ("Compliance", "compliance_rating", "rating"),
    ("CSAT", "csat_rating", "rating"),
    ("Training (TWCS)", "twcs_score", "twcs"),
]
_HEADER_KEYS = {
    "employee id": "code", "emp id": "code", "employee code": "code", "employee name": "name", "name": "name",
    "tasks on-time": "on_time_count", "tasks on time": "on_time_count", "tasks delayed": "delayed_count", "achievement": "achievement_pct",
    "quality": "quality_rating", "collaboration": "collaboration_rating", "innovation": "innovation_rating", "compliance": "compliance_rating",
    "csat": "csat_rating", "training": "twcs_score", "twcs": "twcs_score",
}


def _header_field(h) -> str | None:
    """Match a header by its start, so the original sheet's long headers ('Quality: Manager assessment (1-5 scale)') work."""
    t = " ".join(str(h or "").replace("\n", " ").lower().split())
    for key in sorted(_HEADER_KEYS, key=len, reverse=True):
        if t.startswith(key):
            return _HEADER_KEYS[key]
    return None


@router.get("/months/{month_id}/checkin-template.xlsx")
def checkin_template(month_id: int, ctx: Ctx = Depends(get_ctx)):
    """One row per person the viewer can mark, with the values already entered. Blank cells mean 'use the records'."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    rows = _people_rows(ctx, pilot, month)
    fields = _sheet_fields(ctx, pilot, month)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Check-in {month.month:%b %Y}"
    head = ["Employee ID", "Employee Name", "Team", "Category"] + [h for h, _, _ in fields]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E7F5EB")
    for r in rows:
        if not r["can_edit"]:
            continue
        ci = r["check_in"]
        vals = [ci.get(f) if kind == "rating" else ci["entered"].get(f) for _, f, kind in fields]
        ws.append([r["employee"]["code"], r["employee"]["name"], r["employee"]["department"], pilot.dept.category, *vals])
    for i, w in enumerate([13, 24, 10, 15] + [14] * len(fields), start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    notes = wb.create_sheet("How to fill")
    for line in [
        ["Monthly check-in upload", f"{month.month:%b %Y}"],
        ["Employee ID", "Required. Rows are matched by Employee ID (name is used only if the ID is blank)."],
        ["Tasks On-Time, Tasks Delayed, Achievement %", "Not uploaded: worked out from the KPIs marked in Team."],
        [", ".join(h for h, _, k in fields if k == "rating"), "Whole numbers 1 to 5."],
        ["Training (TWCS)", "Timeliness-Weighted Completion Score, 0 to 100. Blank = worked out from the training modules."],
        ["Blank cells", "Leave the current value unchanged."],
    ]:
        notes.append(line)
    notes.column_dimensions["A"].width = 44
    notes.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    record(ctx.db, ctx.account.id, "download_checkin_template", "kpi_month", month.id, ind.month_label(month.month))
    ctx.db.commit()
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="checkin-{month.month:%Y-%m}.xlsx"'})


def _sheet_fields(ctx: Ctx, pilot: ind.Pilot, month: KpiMonth) -> list[tuple]:
    """Ratings weighted in the department's category, plus Training. Non-weighted ratings are not asked for."""
    weights = ctx.db.get(ScoringRule, month.rule_id).config["weights"].get(pilot.dept.category, {})
    return [(h, f, k) for h, f, k in SHEET_FIELDS if k != "rating" or weights.get(f.removesuffix("_rating"))]


def _cell_value(raw, kind: str):
    """Returns (value, error). Blank -> (None, None) meaning 'no change'."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None, None
    try:
        v = float(str(raw).strip().rstrip("%"))
    except ValueError:
        return None, f"'{raw}' is not a number"
    if kind == "rating":
        if v != int(v) or not 1 <= v <= 5:
            return None, f"{raw} must be a whole number from 1 to 5"
        return int(v), None
    if kind == "count":
        if v != int(v) or v < 0:
            return None, f"{raw} must be a whole number of 0 or more"
        return int(v), None
    if kind == "pct" and not 0 <= v <= 150:
        return None, f"{raw} must be between 0 and 150"
    if kind == "twcs" and not 0 <= v <= 100:
        return None, f"{raw} must be between 0 and 100"
    return v, None


@router.post("/months/{month_id}/checkin-upload")
def checkin_upload(month_id: int, dry_run: bool = True, file: UploadFile = File(...), ctx: Ctx = Depends(get_ctx)):
    """Upload the R&R parameter columns for many people. Dry run first: every row reports what would change or why it cannot."""
    import openpyxl

    pilot = _pilot(ctx)
    month = ind.require_month(ctx.db, month_id)
    _require_open(month)
    data = file.file.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        raise RuleViolation("The file is larger than 5 MB.", "file")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise RuleViolation("Upload an Excel .xlsx file (use the template).", "file")
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header, cols = None, {}
    for row in it:  # the header is the first row that names at least the employee and one parameter
        found = {i: _header_field(h) for i, h in enumerate(row)}
        found = {i: f for i, f in found.items() if f}
        if ("code" in found.values() or "name" in found.values()) and len(found) >= 2:
            header, cols = row, found
            break
    if header is None:
        raise RuleViolation("No header row found. Keep the template's column names (Employee ID, Tasks On-Time, Quality, ...).", "file")
    kinds = {f: k for _, f, k in _sheet_fields(ctx, pilot, month)}
    ignored = sorted({AUTO_FIELDS[f] for f in cols.values() if f in AUTO_FIELDS})
    by_code = {e.code.lower(): e for e in pilot.people}
    by_name: dict[str, list] = {}
    for e in pilot.people:
        by_name.setdefault(" ".join(e.full_name.lower().split()), []).append(e)
    visible = ind.visible_ids(pilot, ctx.scope)
    results, seen = [], set()
    for row_no, row in enumerate(it, start=2):
        vals = {f: row[i] if i < len(row) else None for i, f in cols.items()}
        if all(v is None or str(v).strip() == "" for v in vals.values()):
            continue
        code = str(vals.get("code") or "").strip()
        name = " ".join(str(vals.get("name") or "").lower().split())
        emp = by_code.get(code.lower()) if code else (by_name.get(name, [None])[0] if len(by_name.get(name, [])) == 1 else None)
        res = {"row": row_no, "code": code, "name": vals.get("name"), "employee": person(emp) if emp else None, "changes": {}, "errors": []}
        if emp is None:
            res["errors"].append("No Sales pilot employee matches this Employee ID" + ("" if code else " or name (add the Employee ID)") + ".")
        elif emp.id not in visible:
            res["errors"].append("This person is outside your scope.")
        elif emp.id in seen:
            res["errors"].append("This person appears more than once in the file.")
        else:
            seen.add(emp.id)
            ok, why = ind.can_edit(pilot, ctx.account, emp, month, ind.get_evaluation(ctx.db, month, emp.id))
            if not ok:
                res["errors"].append(why)
        for f, kind in kinds.items():
            if f not in vals:
                continue
            v, err = _cell_value(vals[f], kind)
            if err:
                res["errors"].append(f"{next(h for h, ff, _ in SHEET_FIELDS if ff == f)}: {err}")
            elif v is not None:
                res["changes"][f] = v
        results.append(res)
    applied = 0
    if not dry_run:
        for res in results:
            if res["errors"] or not res["changes"]:
                continue
            emp = pilot.by_id()[res["employee"]["id"]]
            ev = ind.get_evaluation(ctx.db, month, emp.id)
            if ev is None:
                ev = MonthlyEvaluation(month_id=month.id, employee_id=emp.id, state="Draft", version=0)
                ctx.db.add(ev)
            old = {k: getattr(ev, k) for k in res["changes"]}
            for k, v in res["changes"].items():
                setattr(ev, k, v)
            ev.updated_by = ctx.account.id
            ev.version = (ev.version or 0) + 1
            ctx.db.flush()
            record(ctx.db, ctx.account.id, "upload_checkin", "monthly_evaluation", ev.id, f"{emp.full_name} {ind.month_label(month.month)}", old=old, new=res["changes"], reason=f"Excel upload {file.filename}, row {res['row']}")
            applied += 1
        ctx.db.commit()
    return {"dry_run": dry_run, "file": file.filename, "rows": results, "ready": sum(1 for r in results if not r["errors"] and r["changes"]),
            "with_errors": sum(1 for r in results if r["errors"]), "applied": applied,
            "ignored_columns": ignored + sorted({h for h, f, _ in SHEET_FIELDS if f in cols.values() and f not in kinds})}


# --- demo data (Admin) ---------------------------------------------------------------------------


@router.post("/demo/{action}")
def demo_data(action: str, ctx: Ctx = Depends(get_ctx)):
    """Admin: add or remove the labelled Individual Tier demo data. Real values and the audit log are never touched."""
    ctx.scope.require_role("admin")
    if action not in ("add", "clear"):
        raise NotFound("Action")
    out = seed_individual_demo(ctx.db) if action == "add" else clear_individual_demo(ctx.db)
    record(ctx.db, ctx.account.id, f"demo_{action}", "kpi_month", None, "Individual Tier demo data", new=out)
    ctx.db.commit()
    return out


# --- history for one person and month ------------------------------------------------------------

ACTION_TEXT = {
    "create_kpi_item": "Added", "update_kpi_item": "Changed", "remove_kpi_item": "Removed", "resolve_warning": "Checked a warning on",
    "delay_validate": "Validated the delay on", "delay_reject": "Rejected the delay reason on",
    "create_training_module": "Assigned training", "update_training_module": "Updated training", "verify_training_module": "Verified training",
    "request_training_exemption": "Requested an exemption for", "exemption_approve": "Approved the exemption for", "exemption_reject": "Rejected the exemption for",
    "save_evaluation": "Saved the check-in", "upload_checkin": "Uploaded check-in values", "submit_evaluation": "Submitted the check-in",
    "approve_evaluation": "Approved the check-in", "return_evaluation": "Returned the check-in",
}


@router.get("/history/{employee_id}/{month_id}")
def person_history(employee_id: int, month_id: int, ctx: Ctx = Depends(get_ctx)):
    """Who changed what for this person and month, newest first (from the append-only audit log)."""
    pilot = _pilot(ctx)
    emp = ind.require_person(pilot, ctx.scope, employee_id)
    month = ind.require_month(ctx.db, month_id)
    item_ids = [str(i) for i in ctx.db.scalars(select(KpiItem.id).where(KpiItem.month_id == month.id, KpiItem.employee_id == emp.id))]
    start, end = month.month, month_end(month.month)
    module_ids = [str(i) for i in ctx.db.scalars(select(TrainingModule.id).where(TrainingModule.employee_id == emp.id, TrainingModule.due_date >= start, TrainingModule.due_date <= end))]
    ev = ind.get_evaluation(ctx.db, month, emp.id)
    conds = []
    if item_ids:
        conds.append((AuditEvent.entity == "kpi_item") & AuditEvent.entity_id.in_(item_ids))
    if module_ids:
        conds.append((AuditEvent.entity == "training_module") & AuditEvent.entity_id.in_(module_ids))
    if ev is not None:
        conds.append((AuditEvent.entity == "monthly_evaluation") & (AuditEvent.entity_id == str(ev.id)))
    if not conds:
        return []
    from sqlalchemy import or_

    rows = ctx.db.scalars(select(AuditEvent).where(or_(*conds)).order_by(AuditEvent.at.desc(), AuditEvent.id.desc()).limit(200))
    method = ind.method_visible(ctx.account)
    out = []
    for e in rows:
        new = e.new or {}
        if not method:  # never leak scoring internals through history
            new = {k: v for k, v in new.items() if k not in ("coverage", "base", "penalty_total")}
        out.append({"id": e.id, "at": e.at, "actor": account_person(ctx.db, e.actor_account_id), "action": e.action,
                    "text": ACTION_TEXT.get(e.action, e.action.replace("_", " ").capitalize()), "label": e.label, "reason": e.reason,
                    "old": e.old, "new": new})
    return out


# --- corrections to a closed month ---------------------------------------------------------------


def _correction_dict(ctx: Ctx, c: CorrectionRequest) -> dict:
    return {"id": c.id, "month_id": c.month_id, "employee_id": c.employee_id, "reason": c.reason, "state": c.state,
            "requested_by": account_person(ctx.db, c.requested_by), "decided_by": account_person(ctx.db, c.decided_by),
            "decided_at": c.decided_at, "decision_note": c.decision_note, "completed_at": c.completed_at, "created_at": c.created_at,
            "old_result": c.old_result if ind.method_visible(ctx.account) else _public_result(c.old_result),
            "new_result": c.new_result if ind.method_visible(ctx.account) else _public_result(c.new_result),
            "can_decide": c.state == "Requested" and ctx.account.role == "admin" and c.requested_by != ctx.account.id}


def _public_result(r: dict | None) -> dict | None:
    return None if r is None else {k: r.get(k) for k in ("final", "tier", "status")}


class CorrectionIn(BaseModel):
    month_id: int
    employee_id: int
    reason: str = Field(min_length=10)


@router.post("/corrections")
def request_correction(body: CorrectionIn, ctx: Ctx = Depends(get_ctx)):
    """HOD or Admin asks to fix one person's result in a closed month. An Admin other than the requester approves."""
    ctx.scope.require_role("hod", "admin")
    pilot = _pilot(ctx)
    emp = ind.require_person(pilot, ctx.scope, body.employee_id)
    month = ind.require_month(ctx.db, body.month_id)
    if month.state != "Locked":
        raise RuleViolation("Corrections are for closed months. In an open month, return the check-in instead.", "month_id")
    if ctx.db.scalar(select(CorrectionRequest.id).where(CorrectionRequest.month_id == month.id, CorrectionRequest.employee_id == emp.id,
                                                      CorrectionRequest.state.in_(("Requested", "Open")))):
        raise RuleViolation("A correction for this person and month is already in progress.", "employee_id")
    c = CorrectionRequest(month_id=month.id, employee_id=emp.id, reason=body.reason.strip(), requested_by=ctx.account.id)
    ctx.db.add(c)
    ctx.db.flush()
    label = f"{emp.full_name} {ind.month_label(month.month)}"
    record(ctx.db, ctx.account.id, "request_correction", "correction", c.id, label, reason=c.reason)
    for acc in ctx.db.scalars(select(Account).where(Account.active.is_(True), Account.role == "admin", Account.id != ctx.account.id)):
        notify(ctx.db, acc.id, f"Correction to approve: {label}", c.reason, f"/individual/performance/{emp.id}/{month.id}")
    ctx.db.commit()
    return _correction_dict(ctx, c)


@router.post("/corrections/{correction_id}/decide")
def decide_correction(correction_id: int, body: DecisionIn, ctx: Ctx = Depends(get_ctx)):
    """Admin approves (the person's closed month reopens for their Manager) or rejects with a reason."""
    ctx.scope.require_role("admin")
    pilot = _pilot(ctx)
    c = ctx.db.get(CorrectionRequest, correction_id)
    if c is None:
        raise NotFound("Correction")
    emp = ind.require_person(pilot, ctx.scope, c.employee_id)
    month = ind.require_month(ctx.db, c.month_id)
    if c.state != "Requested":
        raise RuleViolation("Only a requested correction can be decided.", "state")
    if c.requested_by == ctx.account.id:
        raise Forbidden("The person who requested a correction cannot approve it.")
    if body.decision not in ("approve", "reject"):
        raise RuleViolation("Choose approve or reject.", "decision")
    if body.decision == "reject" and not body.reason.strip():
        raise RuleViolation("Give the reason for rejecting it.", "reason")
    c.decided_by, c.decided_at, c.decision_note = ctx.account.id, now_ist(), body.reason.strip() or None
    label = f"{emp.full_name} {ind.month_label(month.month)}"
    if body.decision == "approve":
        c.state = "Open"
        ev = ind.get_evaluation(ctx.db, month, emp.id)
        if ev is None:
            ev = MonthlyEvaluation(month_id=month.id, employee_id=emp.id, version=0)
            ctx.db.add(ev)
        ev.state, ev.decision_note, ev.excluded_reason = "Returned", f"Correction approved: {c.reason}", None
        ev.version = (ev.version or 0) + 1
        evaluator = pilot.evaluator(emp)
        if evaluator:
            notify(ctx.db, evaluator.id, f"Correction opened: {label}", c.reason, f"/individual/team/{emp.id}?month={month.id}")
    else:
        c.state = "Rejected"
    record(ctx.db, ctx.account.id, f"{body.decision}_correction", "correction", c.id, label, new={"state": c.state}, reason=body.reason or None)
    notify(ctx.db, c.requested_by, f"Correction {c.state.lower()}: {label}", body.reason or None, f"/individual/performance/{emp.id}/{month.id}")
    ctx.db.commit()
    return _correction_dict(ctx, c)


@router.get("/corrections")
def list_corrections(month_id: int | None = None, employee_id: int | None = None, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    ids = ind.visible_ids(pilot, ctx.scope)
    q = select(CorrectionRequest).order_by(CorrectionRequest.id.desc())
    if month_id:
        q = q.where(CorrectionRequest.month_id == month_id)
    if employee_id:
        q = q.where(CorrectionRequest.employee_id == employee_id)
    return [_correction_dict(ctx, c) for c in ctx.db.scalars(q) if c.employee_id in ids]


# --- nominations (monthly recognition) -----------------------------------------------------------


CITATION_MIN_WORDS = 50


def _nomination_dict(ctx: Ctx, n: MonthlyNomination, pilot: ind.Pilot) -> dict:
    emp = pilot.by_id().get(n.employee_id)
    can_decide = (n.state == "Submitted" and ctx.account.role in ("hod", "admin") and n.nominated_by != ctx.account.id
                  and n.employee_id != ctx.account.employee_id)
    return {"id": n.id, "month_id": n.month_id, "employee": person(emp) if emp else None, "award": n.award, "citation": n.citation,
            "evidence": n.evidence, "state": n.state, "nominated_by": account_person(ctx.db, n.nominated_by),
            "decided_by": account_person(ctx.db, n.decided_by), "decided_at": n.decided_at, "decision_note": n.decision_note,
            "created_at": n.created_at, "illustrative": n.illustrative, "can_decide": can_decide}


@router.get("/nominations")
def list_nominations(month_id: int, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    ids = ind.visible_ids(pilot, ctx.scope)
    month = ind.require_month(ctx.db, month_id)
    rows = ctx.db.scalars(select(MonthlyNomination).where(MonthlyNomination.month_id == month.id).order_by(MonthlyNomination.id))
    # Published approved nominations are visible to everyone in the pilot; the rest stay within scope.
    return {"awards": list(NOMINATION_AWARDS), "min_words": CITATION_MIN_WORDS,
            "rows": [_nomination_dict(ctx, n, pilot) for n in rows if n.employee_id in ids or (month.published_at and n.state == "Approved")]}


class NominationIn(BaseModel):
    month_id: int
    employee_id: int
    award: str
    citation: str
    evidence: str | None = None


@router.post("/nominations")
def nominate(body: NominationIn, ctx: Ctx = Depends(get_ctx)):
    """Nominate someone from a closed month. A citation of at least 50 words is required; someone else decides."""
    pilot = _pilot(ctx)
    emp = ind.require_person(pilot, ctx.scope, body.employee_id)
    month = ind.require_month(ctx.db, body.month_id)
    if month.state != "Locked":
        raise RuleViolation("Nominations open once the month is closed.", "month_id")
    if body.award not in NOMINATION_AWARDS:
        raise RuleViolation("Choose an award.", "award")
    words = len(body.citation.split())
    if words < CITATION_MIN_WORDS:
        raise RuleViolation(f"Write at least {CITATION_MIN_WORDS} words in the citation ({words} so far).", "citation")
    if ctx.db.scalar(select(MonthlyNomination.id).where(MonthlyNomination.month_id == month.id, MonthlyNomination.employee_id == emp.id,
                                                       MonthlyNomination.award == body.award, MonthlyNomination.state != "Rejected")):
        raise RuleViolation("This person already has this nomination for the month.", "award")
    n = MonthlyNomination(month_id=month.id, employee_id=emp.id, award=body.award, citation=body.citation.strip(),
                          evidence=(body.evidence or "").strip() or None, nominated_by=ctx.account.id)
    ctx.db.add(n)
    ctx.db.flush()
    label = f"{body.award}: {emp.full_name} {ind.month_label(month.month)}"
    record(ctx.db, ctx.account.id, "nominate", "nomination", n.id, label, new={"award": body.award})
    for acc in ctx.db.scalars(select(Account).where(Account.active.is_(True), Account.role.in_(("hod", "admin")), Account.id != ctx.account.id)):
        if acc.role == "admin" or acc.employee_id in {h.id for h in pilot.heads}:
            notify(ctx.db, acc.id, f"Nomination to decide: {label}", None, f"/individual/rewards?month={month.id}")
    ctx.db.commit()
    return _nomination_dict(ctx, n, pilot)


@router.post("/nominations/{nomination_id}/decide")
def decide_nomination(nomination_id: int, body: DecisionIn, ctx: Ctx = Depends(get_ctx)):
    pilot = _pilot(ctx)
    n = ctx.db.get(MonthlyNomination, nomination_id)
    if n is None:
        raise NotFound("Nomination")
    emp = ind.require_person(pilot, ctx.scope, n.employee_id)
    d = _nomination_dict(ctx, n, pilot)
    if not d["can_decide"]:
        raise Forbidden("Nominations are decided by the HOD or an Admin, never by the nominator or the person nominated.")
    if body.decision not in ("approve", "reject"):
        raise RuleViolation("Choose approve or reject.", "decision")
    if body.decision == "reject" and not body.reason.strip():
        raise RuleViolation("Give the reason for rejecting it.", "reason")
    n.state = "Approved" if body.decision == "approve" else "Rejected"
    n.decided_by, n.decided_at, n.decision_note = ctx.account.id, now_ist(), body.reason.strip() or None
    month = ctx.db.get(KpiMonth, n.month_id)
    label = f"{n.award}: {emp.full_name} {ind.month_label(month.month)}"
    record(ctx.db, ctx.account.id, f"{body.decision}_nomination", "nomination", n.id, label, reason=body.reason or None)
    notify(ctx.db, n.nominated_by, f"Nomination {n.state.lower()}: {label}", body.reason or None, f"/individual/rewards?month={month.id}")
    ctx.db.commit()
    return _nomination_dict(ctx, n, pilot)


@router.get("/search")
def search_people(q: str, ctx: Ctx = Depends(get_ctx)):
    """Individual Tier search: people in the viewer's pilot scope, by name or employee ID."""
    pilot = _pilot(ctx)
    term = " ".join(q.lower().split())
    if len(term) < 2:
        return {"people": []}
    ids = ind.visible_ids(pilot, ctx.scope)
    hits = [e for e in pilot.people if e.id in ids and (term in e.full_name.lower() or term in e.code.lower())]
    return {"people": [person(e) for e in sorted(hits, key=lambda e: e.full_name)[:12]]}
