"""Import of PT_RR_Corrected_Data.xlsx (canonical seed/migration data).

Rules applied (CLAUDE.md section 16):
- Positions are separate from employees; hidden outdated org sheets were already excluded by the workbook.
- Tracker statuses use the workbook's Mapped Status (Ongoing -> In Progress, Delayed -> Blocked/In Progress,
  Completed without date -> In Review). Completed-with-date rows become legacy accepted work (source='import').
- Source defects (two Owners, shares != 100%, unresolved names) are imported with flags, never silently fixed.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models import (
    Assignment,
    Blocker,
    Department,
    DepartmentHead,
    Employee,
    Project,
    ReportingLine,
    ReviewDecision,
    Risk,
    Submission,
    Task,
    TemplateStage,
    WorkflowTemplate,
)

ORG_CHART_DATE = date(2026, 8, 23)

BLOCKER_MAP = {
    "Management reprioritisation": "Management reprioritisation",
    "Internal dependency (Product/Development)": "Internal dependency",
    "Internal dependency (other team review)": "Internal dependency",
    "Requirement change": "Requirement change",
    "Technical issue": "Technical issue",
    "Rework (own team feedback)": "Own-team rework",
}


def _rows(wb, sheet: str) -> list[dict]:
    it = wb[sheet].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    out = []
    for r in it:
        if any(v not in (None, "") for v in r):
            out.append(dict(zip(header, r)))
    return out


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _num(v) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load_workbook(path: Path | None = None):
    return openpyxl.load_workbook(path or settings.seed_workbook, read_only=True, data_only=True)


def import_workbook(db: Session, path: Path | None = None) -> dict:
    wb = load_workbook(path)
    counts: dict[str, int] = defaultdict(int)

    # Departments -------------------------------------------------------------------------------
    depts: dict[str, Department] = {}
    head_codes: dict[str, str] = {}
    for r in _rows(wb, "Departments"):
        name, cat = r.get("Department"), r.get("Proposed Category")
        if not name or not cat or str(name).startswith(("Total", "Note")):
            continue
        proposed = "(new)" in str(cat)
        d = Department(name=str(name).strip(), category=str(cat).replace("(new)", "").strip(), category_proposed=proposed, aliases=r.get("Tracker Name Aliases"), notes=r.get("Notes"))
        db.add(d)
        depts[d.name] = d
        if r.get("Head Record ID"):
            head_codes[d.name] = str(r["Head Record ID"]).strip()
        counts["departments"] += 1
    db.flush()

    # Employees and positions -------------------------------------------------------------------
    emps: dict[str, Employee] = {}
    manager_codes: dict[str, str] = {}
    for r in _rows(wb, "Employee_Master"):
        code = str(r.get("Record ID") or "").strip()
        if not code or code in emps:
            continue  # heads appear once per headed department; one identity
        is_position = str(r.get("Record Type") or "").lower().startswith("vacant")
        name = str(r.get("Full Name") or "").strip()
        if is_position or name in ("", "(vacant)"):
            name = f"{r.get('Designation') or 'Position'} (vacant)"
        dept = depts.get(str(r.get("Department") or "").strip())
        e = Employee(
            code=code,
            hrms_code=r.get("HRMS Emp Code (fill)"),
            record_type="position" if is_position else "employee",
            full_name=name,
            designation=r.get("Designation"),
            department_id=dept.id if dept else None,
            level=int(r["Level in Chart"]) if _num(r.get("Level in Chart")) is not None else None,
            proposed_role=r.get("Proposed Platform Role"),
            flags=r.get("Flags / To Confirm"),
        )
        db.add(e)
        emps[code] = e
        if r.get("Reporting Manager ID"):
            manager_codes[code] = str(r["Reporting Manager ID"]).strip()
        counts["employees" if not is_position else "positions"] += 1
    db.flush()
    for code, mgr in manager_codes.items():
        if mgr in emps:
            emps[code].manager_id = emps[mgr].id
            db.add(ReportingLine(employee_id=emps[code].id, manager_id=emps[mgr].id, department_id=emps[code].department_id, effective_from=ORG_CHART_DATE, reason="Org chart 23-Aug-26 (imported)"))
    for dname, hcode in head_codes.items():
        if hcode in emps:
            db.add(DepartmentHead(department_id=depts[dname].id, employee_id=emps[hcode].id))
    db.flush()

    def resolve_owner(name: str | None, dept: Department | None) -> Employee | None:
        if not name:
            return None
        name = str(name).strip()
        cands = [e for e in emps.values() if e.record_type == "employee" and (e.full_name == name or e.full_name.startswith(name + " "))]
        if dept:
            in_dept = [e for e in cands if e.department_id == dept.id]
            if len(in_dept) == 1:
                return in_dept[0]
            heads = {emps[c].id for d, c in head_codes.items() if depts[d].id == dept.id and c in emps}
            headed = [e for e in cands if e.id in heads]
            if len(headed) == 1:
                return headed[0]
        return cands[0] if len(cands) == 1 else None

    # Projects ----------------------------------------------------------------------------------
    projects: dict[str, Project] = {}
    for r in _rows(wb, "Project_Register"):
        pid = str(r.get("Project ID") or "")
        if not pid.startswith("PRJ"):
            continue
        dept = depts.get(str(r.get("Department") or "").strip())
        owner = resolve_owner(r.get("Owner"), dept)
        p = Project(code=pid, title=str(r["Project"]).strip(), department_id=dept.id if dept else None, owner_id=owner.id if owner else None, priority=r.get("Priority") or None, status="In Progress", source="import")
        db.add(p)
        projects[pid] = p
        counts["projects"] += 1
        if not r.get("Priority"):
            counts["projects_missing_priority"] += 1
    db.flush()

    # Tasks and assignments ---------------------------------------------------------------------
    by_task: dict[str, list[dict]] = defaultdict(list)
    for r in _rows(wb, "Tracker_Cleaned"):
        if r.get("Task ID"):
            by_task[str(r["Task ID"])].append(r)
    for task_code, rows in by_task.items():
        r0 = rows[0]
        p = projects.get(str(r0["Project ID"]))
        if p is None:
            continue
        status = r0.get("Mapped Status") or "Not Started"
        completed = _date(r0.get("Completed Date"))
        due = _date(r0.get("Target Date"))
        flags = [str(r0["Data Quality Flags"])] if r0.get("Data Quality Flags") else []
        roles = [x.get("Assignment Role") for x in rows if x.get("Emp ID")]
        share_total = sum((_num(x.get("Contribution %")) or 0) for x in rows if x.get("Emp ID"))
        if roles.count("Owner") != 1:
            flags.append(f"Source has {roles.count('Owner')} Owner(s); exactly one required")
        if round(share_total * 100) != 100:
            flags.append(f"Source shares total {share_total * 100:g}%")
        unresolved = [str(x.get("Assignee (as typed)")) for x in rows if not x.get("Emp ID")]
        if unresolved:
            flags.append("Unresolved assignee: " + ", ".join(unresolved))
        t = Task(
            code=task_code,
            project_id=p.id,
            title=str(r0.get("Task / Deliverable") or p.title),
            description=r0.get("Delay Reason"),
            complexity=3,
            start_date=_date(r0.get("Start Date")),
            due_date_original=due,
            due_date=due,
            status=status,
            progress=100 if status in ("Completed", "In Review") else 0,
            source="import",
            data_flags="; ".join(flags) or None,
        )
        if status == "Completed" and completed:
            t.accepted_submission_date = completed
            t.accepted_at = datetime.combine(completed, datetime.min.time())
        db.add(t)
        db.flush()
        for x in rows:
            emp = emps.get(str(x.get("Emp ID") or ""))
            if emp:
                db.add(Assignment(task_id=t.id, employee_id=emp.id, role=x.get("Assignment Role") or "Contributor", share=_num(x.get("Contribution %")) or 0))
                counts["assignments"] += 1
        if status == "Completed" and completed:
            s = Submission(task_id=t.id, version=1, submission_date=completed, note="Legacy tracker completion (migration)", source="import")
            db.add(s)
            db.flush()
            db.add(ReviewDecision(task_id=t.id, submission_id=s.id, decision="accept", source="import", feedback="Imported as accepted: Completed with a completion date in the corrected workbook."))
        btype = r0.get("Blocker Type")
        if btype in BLOCKER_MAP:
            requester = emps.get(str(r0.get("Emp ID") or ""))
            db.add(
                Blocker(
                    task_id=t.id,
                    category=BLOCKER_MAP[btype],
                    description=r0.get("Delay Reason"),
                    requester_id=requester.id if requester else None,
                    state="Validated" if r0.get("Validated?") == "Yes" else "Pending",
                    decision_reason="Validated in source tracker (migration)" if r0.get("Validated?") == "Yes" else None,
                    source="import",
                )
            )
            counts["blockers"] += 1
        counts["tasks"] += 1
    # derive project dates/status from tasks
    for p in projects.values():
        ts = [t for t in db.query(Task).filter(Task.project_id == p.id)]
        starts = [t.start_date for t in ts if t.start_date]
        dues = [t.due_date for t in ts if t.due_date]
        p.start_date = min(starts) if starts else None
        p.target_date = max(dues) if dues else None
        statuses = {t.status for t in ts}
        if statuses and statuses <= {"Completed"}:
            p.status = "Completed"
        elif statuses and statuses <= {"On Hold", "Completed"}:
            p.status = "On Hold"
        elif "Blocked" in statuses and not statuses & {"In Progress", "Not Started"}:
            p.status = "Blocked"
        elif statuses <= {"Not Started"}:
            p.status = "Not Started"
        else:
            p.status = "In Progress"

    # Workflow template + risks -----------------------------------------------------------------
    tpl = WorkflowTemplate(name="Email campaign", description="From RACI.xlsx (named Sheet3 defaults). Source conflicts must be resolved before activation.", state="Draft")
    db.add(tpl)
    db.flush()
    for r in _rows(wb, "Workflow_Templates"):
        if r.get("Template") == "Email campaign" and isinstance(r.get("Step"), (int, float)):
            resp = emps.get(str(r.get("R Emp ID") or "").strip())
            acc = emps.get(str(r.get("A Emp ID") or "").strip())
            reviewer_text = str(r.get("Reviewer (validation owner)") or "")
            reviewer = resp if resp and resp.full_name.split()[0] in reviewer_text and "Proposed" not in reviewer_text else None
            db.add(
                TemplateStage(
                    template_id=tpl.id,
                    step=int(r["Step"]),
                    name=str(r.get("Stage (task)")),
                    depends_on_step=int(r["Depends on step"]) if _num(r.get("Depends on step")) else None,
                    responsible_role=r.get("Responsible (R) role"),
                    responsible_id=resp.id if resp else None,
                    accountable_label=r.get("Accountable (A) - exactly one"),
                    accountable_id=acc.id if acc else None,
                    consulted=r.get("Consulted (C)"),
                    informed=r.get("Informed (I)"),
                    checklist=r.get("Checklist (acceptance criteria)"),
                    reviewer_label=reviewer_text,
                    reviewer_id=reviewer.id if reviewer else None,
                    review_criteria=r.get("Review criteria"),
                    required_evidence=r.get("Required evidence before submit"),
                    sla_target_days=_num(r.get("SLA target (business days)")),
                    sla_max_days=_num(r.get("SLA max (business days)")),
                    effort_min_hours=_num(r.get("Effort min (hrs)")),
                    effort_max_hours=_num(r.get("Effort max (hrs)")),
                    mitigation=r.get("Mitigation"),
                    source_flags=r.get("Flags / to confirm"),
                )
            )
            counts["template_stages"] += 1
        elif str(r.get("Template") or "").startswith("End-to-end"):
            tpl.sla_target_days, tpl.sla_max_days = _num(r.get("Step")), _num(r.get("Stage (task)"))
            tpl.effort_min_hours, tpl.effort_max_hours = _num(r.get("Depends on step")), _num(r.get("Responsible (R) role"))
    # Fallback to the documented totals if the summary row layout differs.
    tpl.sla_target_days = tpl.sla_target_days or 9
    tpl.sla_max_days = tpl.sla_max_days or 14
    tpl.effort_min_hours = tpl.effort_min_hours or 15.5
    tpl.effort_max_hours = tpl.effort_max_hours or 25

    cs = WorkflowTemplate(name="Client Support issue flow", description="Recurring Client Support pattern from the tracker. Stage SLAs and named people are not defined in the sources.", state="Draft")
    db.add(cs)
    db.flush()
    for i, name in enumerate(["Design and Testing", "Co-ordination with Product", "Implementation sign-off"], start=1):
        db.add(TemplateStage(template_id=cs.id, step=i, name=name, depends_on_step=i - 1 or None, responsible_role="Client Support team", review_criteria="Acceptance criteria to be defined", source_flags="People and SLA not defined in sources"))

    for r in _rows(wb, "Risk_Register"):
        if str(r.get("Risk ID") or "").startswith("RSK"):
            db.add(Risk(code=r["Risk ID"], template_id=tpl.id, stage=r.get("Stage"), risk=r["Risk"], criticality=r.get("Criticality") or "Medium", dependencies=r.get("Dependencies"), assumptions=r.get("Assumptions"), mitigation=r.get("Mitigation"), state=r.get("Status") or "Open"))
            counts["template_risks"] += 1
    db.flush()
    return dict(counts)


def dry_run_report(path: Path | None = None) -> dict:
    """Severity-coded validation of the workbook without writing anything."""
    wb = load_workbook(path)
    emp = _rows(wb, "Employee_Master")
    tracker = _rows(wb, "Tracker_Cleaned")
    projects = [r for r in _rows(wb, "Project_Register") if str(r.get("Project ID") or "").startswith("PRJ")]
    issues_src = _rows(wb, "Issue_Register")
    collisions = _rows(wb, "Name_Collisions")
    issues = []
    for r in projects:
        if not r.get("Priority"):
            issues.append({"severity": "Critical", "source": f"Project_Register / {r['Project ID']}", "issue": "Priority missing", "correction": "Choose High, Medium or Low"})
    for r in collisions:
        if r.get("Name") and not r.get("Resolution (fill)"):
            issues.append({"severity": "High", "source": f"Employee_Master / {r.get('Record ID')}", "issue": f"Same name, different identity: {r['Name']}", "correction": "Confirm HRMS ID and department"})
    by_task = defaultdict(list)
    for r in tracker:
        if r.get("Task ID"):
            by_task[r["Task ID"]].append(r)
    for code, rows in by_task.items():
        owners = sum(1 for r in rows if r.get("Assignment Role") == "Owner")
        if owners != 1:
            issues.append({"severity": "High", "source": f"Tracker_Cleaned / {code}", "issue": f"{owners} Owner rows", "correction": "Exactly one Owner; shares total 100%"})
        if rows[0].get("Original Status") == "Completed" and not rows[0].get("Completed Date"):
            issues.append({"severity": "Medium", "source": f"Tracker_Cleaned / {code}", "issue": "Completed without completion date", "correction": "Import as In review; seek evidence"})
    issues.append({"severity": "High", "source": "Workflow_Templates / Email campaign", "issue": "Reviewer is an assignee or unresolved on several stages", "correction": "Select independent named reviewers"})
    return {
        "source": settings.seed_workbook.name,
        "state": "Dry run / nothing imported",
        "counts": {
            "employee_rows": sum(1 for r in emp if r.get("Record ID")),
            "projects": len(projects),
            "assignments": sum(1 for r in tracker if r.get("Task ID")),
            "missing_priority": sum(1 for r in projects if not r.get("Priority")),
            "source_issues": sum(1 for r in issues_src if isinstance(r.get("#"), (int, float))),
            "on_time_eligible": sum(1 for r in tracker if r.get("On-Time Eligible (auto)") == "Yes"),
        },
        "issues": issues,
        "notes": "Hidden outdated Product sheets are excluded; vacant positions stay separate from employees.",
    }
