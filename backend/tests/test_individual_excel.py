"""The monthly check-in workbook: download, mark every sheet, dry run, apply, and the same result as marking by hand.

Runs after test_individual_api (Sep 2026 is locked by then) and uses Oct 2026 with people no other test marks there:
Sakshi Sharma (EMP-0281), Alishba Khan (EMP-0282), Ramandeep Kaur (EMP-0283), Rajat Kapoor (EMP-0285) and Anshul Thakur
(EMP-0287), all evaluated by Rajesh Kumar. No rows are added to October, so the carry-over checks in test_month_flow hold."""

import io
from datetime import datetime

import openpyxl
import pytest

from app.services import checkin_excel as cx

SHEETS = [cx.SHEET_CHECKINS, cx.SHEET_KPIS, cx.SHEET_DUTIES, cx.SHEET_TRAINING, cx.SHEET_HELP]


@pytest.fixture(scope="module")
def months(client, auth):
    return {m["label"]: m for m in client.get("/api/individual/context", headers=auth("sales_hod")).json()["months"]}


def _emp(client, h, month, code):
    rows = client.get(f"/api/individual/months/{month['id']}/people", headers=h).json()["rows"]
    return next(r for r in rows if r["employee"]["code"] == code)


def _items(client, h, month, emp_id):
    return client.get(f"/api/individual/months/{month['id']}/items?employee_id={emp_id}", headers=h).json()


def _modules(client, h, month, emp_id):
    return client.get(f"/api/individual/months/{month['id']}/training?employee_id={emp_id}", headers=h).json()["people"][0]["modules"]


def _download(client, h, month) -> openpyxl.Workbook:
    r = client.get(f"/api/individual/months/{month['id']}/checkin-template.xlsx", headers=h)
    assert r.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(r.content))


def _upload(client, h, month, wb, dry_run: bool):
    buf = io.BytesIO()
    wb.save(buf)
    files = {"file": ("team-checkin.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    return client.post(f"/api/individual/months/{month['id']}/checkin-upload?dry_run={str(dry_run).lower()}", files=files, headers=h)


def _rows(ws):
    """{(Employee ID, Row ID or None): Excel row number} plus the header -> column index map."""
    head = {c.value: c.column for c in ws[cx.HEADER_ROW] if c.value}
    out = {}
    for n in range(cx.HEADER_ROW + 1, ws.max_row + 1):
        code = ws.cell(n, head["Employee ID"]).value
        if code:
            out[(code, ws.cell(n, head[cx.ROW_ID]).value if cx.ROW_ID in head else None)] = n
    return head, out


def _state(client, h, month, emp_id):
    """Everything marked for one person, and the result, as the person page sees it."""
    items = _items(client, h, month, emp_id)
    ev = client.get(f"/api/individual/evaluations/{emp_id}/{month['id']}", headers=h).json()["evaluation"] or {}
    row = next(r for r in client.get(f"/api/individual/months/{month['id']}/people", headers=h).json()["rows"] if r["employee"]["id"] == emp_id)
    return {
        "items": sorted((i["id"], i["achieved_value"], i["outcome"], i["completed_date"], i["delay_reason"]) for i in items),
        "modules": sorted((m["id"], m["completed_date"]) for m in _modules(client, h, month, emp_id)),
        "ratings": (ev.get("quality_rating"), ev.get("csat_rating"), ev.get("comments")),
        "result": (row["achievement_pct"], row["on_time"], row["twcs"], row["final"], row["tier"], row["status"]),
    }


def test_template_covers_everything_marked_on_the_person_page(client, auth, months):
    rajesh = auth("sales_rajesh")
    oct_ = months["Oct 2026"]
    alishba = _emp(client, rajesh, oct_, "EMP-0282")["employee"]["id"]
    # Alishba's check-in is submitted, so her rows are read-only in the file.
    assert client.put(f"/api/individual/evaluations/{alishba}/{oct_['id']}", json={"quality_rating": 4, "csat_rating": 4}, headers=rajesh).status_code == 200
    assert client.post(f"/api/individual/evaluations/{alishba}/{oct_['id']}/submit", headers=rajesh).status_code == 200

    wb = _download(client, rajesh, oct_)
    assert wb.sheetnames == SHEETS
    ci = wb[cx.SHEET_CHECKINS]
    head, rows = _rows(ci)
    assert list(head) == ["Employee ID", "Name", "Designation", "Manager", "Check-in", "Quality", "CSAT", "Training (TWCS)", "Comments", cx.VERSION]
    assert len(rows) == 12  # everyone Rajesh evaluates; Managers never see their own row
    assert "Finvasia Elevate" in ci["A1"].value and "Oct 2026" in ci["A1"].value and "Rajesh Kumar's team" in ci["A2"].value
    assert ci.freeze_panes == "C5" and ci.protection.sheet
    sakshi_row, alishba_row = rows[("EMP-0281", None)], rows[("EMP-0282", None)]
    q = ci.cell(sakshi_row, head["Quality"])
    assert q.protection.locked is False and q.fill.fgColor.rgb.endswith(cx.EDIT_FILL)
    assert ci.cell(alishba_row, head["Quality"]).protection.locked and ci.cell(alishba_row, head["Check-in"]).value == "Submitted"
    assert ci.cell(alishba_row, 1).comment is not None  # why the row is read-only
    assert ci.column_dimensions["J"].hidden  # Version
    assert any(dv.formula1 == '"1,2,3,4,5"' for dv in ci.data_validations.dataValidation)

    kp = wb[cx.SHEET_KPIS]
    head, rows = _rows(kp)
    assert list(head) == ["Employee ID", "Name", "KPI", "KPI type", "Target", "Achieved", "% of target", "Result", cx.ROW_ID, cx.VERSION]
    n = next(v for (code, _), v in rows.items() if code == "EMP-0281")
    assert kp.cell(n, head["% of target"]).value.startswith("=IF(") and "Under achieved" in kp.cell(n, head["Result"]).value
    assert kp.cell(n, head["Achieved"]).protection.locked is False and kp.cell(n, head["Target"]).protection.locked
    assert kp.cell(n, head["Target"]).number_format == cx.NUM_FMT  # Indian grouping
    assert kp.column_dimensions["I"].hidden and kp.column_dimensions["J"].hidden
    assert {dv.type for dv in kp.data_validations.dataValidation} == {"decimal"}  # Achieved >= 0

    du = wb[cx.SHEET_DUTIES]
    head, _ = _rows(du)
    assert list(head) == ["Employee ID", "Name", "Task", "Due date", "Done?", "Completed on", "Delay reason", cx.ROW_ID, cx.VERSION]
    assert {dv.type for dv in du.data_validations.dataValidation} == {"list", "date"}
    tr = wb[cx.SHEET_TRAINING]
    assert [c.value for c in tr[cx.HEADER_ROW]] == ["Employee ID", "Name", "Module", "Track", "Due date", "Completed on", cx.ROW_ID, cx.VERSION]
    help_text = " ".join(str(c.value) for row in wb[cx.SHEET_HELP].iter_rows() for c in row if c.value)
    assert "Blank cells" in help_text and "1 (well below" in help_text and "read-only" in help_text.lower()


def test_mark_the_whole_team_in_one_file(client, auth, months):
    rajesh = auth("sales_rajesh")
    oct_ = months["Oct 2026"]
    sakshi = _emp(client, rajesh, oct_, "EMP-0281")["employee"]["id"]
    anshul = _emp(client, rajesh, oct_, "EMP-0287")["employee"]["id"]
    # Sakshi has three KPIs and two duty tasks in October; give her a module due this month too, so every sheet has a row of hers.
    assert sum(1 for i in _items(client, rajesh, oct_, sakshi) if i["kind"] == "duty") == 2
    body = {"employee_id": sakshi, "title": "Excel test module", "track": "Functional", "assigned_date": "2026-10-01", "due_date": "2026-10-10"}
    module_id = client.post("/api/individual/training", json=body, headers=rajesh).json()["id"]
    before = _state(client, rajesh, oct_, sakshi)

    wb = _download(client, rajesh, oct_)
    # Someone changes one of Anshul's KPIs after the download: that row must not be overwritten from the file.
    anshul_kpis = [i for i in _items(client, rajesh, oct_, anshul) if i["kind"] == "kpi"]
    stale = anshul_kpis[1]
    assert client.patch(f"/api/individual/items/{stale['id']}", json={"version": stale["version"], "achieved_value": 1234}, headers=rajesh).status_code == 200

    # --- fill every sheet
    ci = wb[cx.SHEET_CHECKINS]
    head, rows = _rows(ci)
    r = rows[("EMP-0281", None)]
    ci.cell(r, head["Quality"]).value, ci.cell(r, head["CSAT"]).value = 5, 3
    ci.cell(r, head["Comments"]).value = "Strong month on account openings"
    ci.cell(rows[("EMP-0282", None)], head["Quality"]).value = 2  # Alishba is submitted: refused
    gouri = ci.max_row + 1  # Amit's person: outside Rajesh's scope
    ci.cell(gouri, 1).value, ci.cell(gouri, 2).value, ci.cell(gouri, head["Quality"]).value = "EMP-0268", "Gouri", 4

    kp = wb[cx.SHEET_KPIS]
    head, rows = _rows(kp)
    mine = [(rid, n) for (code, rid), n in rows.items() if code == "EMP-0281"]
    targets = {rid: kp.cell(n, head["Target"]).value for rid, n in mine}
    wanted = {}
    for (rid, n), factor in zip(sorted(mine), (0.5, 1.0, 1.25)):  # under, exactly, over
        wanted[rid] = round(targets[rid] * factor)
        kp.cell(n, head["Achieved"]).value = wanted[rid]
    kp.cell(rows[("EMP-0287", stale["id"])], head["Achieved"]).value = 999  # changed since the download
    tampered = rows[("EMP-0287", anshul_kpis[0]["id"])]
    kp.cell(tampered, head[cx.ROW_ID]).value = 999999
    kp.cell(tampered, head["Achieved"]).value = 1

    du = wb[cx.SHEET_DUTIES]
    head, rows = _rows(du)
    (d1, n1), (d2, n2) = sorted((rid, n) for (code, rid), n in rows.items() if code == "EMP-0281")
    du.cell(n1, head["Done?"]).value, du.cell(n1, head["Completed on"]).value = "Completed", datetime(2026, 10, 8)
    du.cell(n2, head["Done?"]).value, du.cell(n2, head["Delay reason"]).value = "Not completed", "Waiting on the dialler vendor"
    ramandeep_duty = next(n for (code, _), n in rows.items() if code == "EMP-0283")
    du.cell(ramandeep_duty, head["Done?"]).value, du.cell(ramandeep_duty, head["Completed on"]).value = "Completed", "09 Jan 2099"  # future

    tr = wb[cx.SHEET_TRAINING]
    head, rows = _rows(tr)
    tr.cell(rows[("EMP-0281", module_id)], head["Completed on"]).value = "12 Oct 2026"  # two days late: 75% credit

    # --- dry run: a per-row preview, nothing saved
    dry = _upload(client, rajesh, oct_, wb, dry_run=True)
    assert dry.status_code == 200, dry.text
    dry = dry.json()
    ready = [r for r in dry["rows"] if not r["errors"]]
    assert {(r["sheet"], r["code"]) for r in ready} == {(cx.SHEET_CHECKINS, "EMP-0281"), (cx.SHEET_KPIS, "EMP-0281"), (cx.SHEET_DUTIES, "EMP-0281"), (cx.SHEET_TRAINING, "EMP-0281")}
    assert dry["ready"] == 7 and dry["with_errors"] == 5 and dry["applied"] == 0
    errors = {(r["sheet"], r["code"]): " ".join(r["errors"]) for r in dry["rows"] if r["errors"]}
    assert "Submitted" in errors[(cx.SHEET_CHECKINS, "EMP-0282")]
    assert "outside your scope" in errors[(cx.SHEET_CHECKINS, "EMP-0268")]
    gouri_row = next(r for r in dry["rows"] if r["code"] == "EMP-0268")
    assert gouri_row["employee"] is None  # no details of people outside the scope
    assert "future" in errors[(cx.SHEET_DUTIES, "EMP-0283")]
    anshul_kpi_errors = [" ".join(r["errors"]) for r in dry["rows"] if r["sheet"] == cx.SHEET_KPIS and r["code"] == "EMP-0287"]
    assert any("Row ID" in e for e in anshul_kpi_errors) and any("after the file was downloaded" in e for e in anshul_kpi_errors)
    checkin = next(r for r in ready if r["sheet"] == cx.SHEET_CHECKINS)
    assert checkin["changes"] == {"quality_rating": 5, "csat_rating": 3, "comments": "Strong month on account openings"} and checkin["old"]["quality_rating"] is None
    assert _state(client, rajesh, oct_, sakshi) == before
    assert next(i for i in _items(client, rajesh, oct_, anshul) if i["id"] == stale["id"])["achieved_value"] == 1234

    # --- apply
    done = _upload(client, rajesh, oct_, wb, dry_run=False).json()
    assert done["applied"] == 7 and done["with_errors"] == 5
    by_file = _state(client, rajesh, oct_, sakshi)
    items = {i["id"]: i for i in _items(client, rajesh, oct_, sakshi)}
    assert {rid: items[rid]["achieved_value"] for rid in wanted} == wanted
    assert [items[rid]["result"] for rid in sorted(wanted)] == ["Under achieved", "Achieved", "Over achieved"]
    assert (items[d1]["outcome"], items[d1]["completed_date"], items[d1]["timing"]) == ("Completed", "2026-10-08", "On time")
    assert (items[d2]["outcome"], items[d2]["delay_reason"], items[d2]["timing"]) == ("Not completed", "Waiting on the dialler vendor", "Delayed")
    assert dict(by_file["modules"])[module_id] == "2026-10-12"
    assert by_file["ratings"] == (5, 3, "Strong month on account openings")
    assert next(i for i in _items(client, rajesh, oct_, anshul) if i["id"] == stale["id"])["achieved_value"] == 1234  # untouched
    hist = client.get(f"/api/individual/history/{sakshi}/{oct_['id']}", headers=rajesh).json()
    assert sum(1 for h in hist if (h["reason"] or "").startswith("Excel upload team-checkin.xlsx")) == 7

    # --- the same values marked by hand on the person page give the same records and the same score
    for rid, v in ((rid, before_v) for rid, before_v, *_ in before["items"]):
        it = items[rid]
        body = {"version": it["version"], "achieved_value": v} if it["kind"] == "kpi" else {"version": it["version"], "outcome": None, "delay_reason": None}
        assert client.patch(f"/api/individual/items/{rid}", json=body, headers=rajesh).status_code == 200
    assert client.patch(f"/api/individual/training/{module_id}", json={"completed_date": None}, headers=rajesh).status_code == 200
    assert client.put(f"/api/individual/evaluations/{sakshi}/{oct_['id']}", json={"quality_rating": None, "csat_rating": None, "comments": None}, headers=rajesh).status_code == 200
    assert _state(client, rajesh, oct_, sakshi)["result"] == before["result"]

    items = {i["id"]: i for i in _items(client, rajesh, oct_, sakshi)}
    for rid, v in wanted.items():
        assert client.patch(f"/api/individual/items/{rid}", json={"version": items[rid]["version"], "achieved_value": v}, headers=rajesh).status_code == 200
    assert client.patch(f"/api/individual/items/{d1}", json={"version": items[d1]["version"], "outcome": "Completed", "completed_date": "2026-10-08"}, headers=rajesh).status_code == 200
    assert client.patch(f"/api/individual/items/{d2}", json={"version": items[d2]["version"], "outcome": "Not completed", "delay_reason": "Waiting on the dialler vendor"}, headers=rajesh).status_code == 200
    assert client.patch(f"/api/individual/training/{module_id}", json={"completed_date": "2026-10-12"}, headers=rajesh).status_code == 200
    body = {"quality_rating": 5, "csat_rating": 3, "comments": "Strong month on account openings"}
    assert client.put(f"/api/individual/evaluations/{sakshi}/{oct_['id']}", json=body, headers=rajesh).status_code == 200
    by_hand = _state(client, rajesh, oct_, sakshi)
    assert by_hand == by_file and by_hand["result"][3] is not None

    # Put the KPI values back: test_month_flow checks October's carried-over KPIs are still unmarked.
    restore = {rid: v for rid, v, *_ in before["items"] if rid in wanted} | {stale["id"]: stale["achieved_value"]}
    for i in _items(client, rajesh, oct_, sakshi) + _items(client, rajesh, oct_, anshul):
        if i["id"] in restore:
            assert client.patch(f"/api/individual/items/{i['id']}", json={"version": i["version"], "achieved_value": restore[i["id"]]}, headers=rajesh).status_code == 200


def test_upload_keeps_the_person_page_rules(client, auth, months):
    rajesh, amit = auth("sales_rajesh"), auth("sales_amit")
    oct_ = months["Oct 2026"]
    wb = _download(client, rajesh, oct_)
    # KPIs have no completion date or delay, even if someone adds those columns to the KPI sheet.
    kp = wb[cx.SHEET_KPIS]
    head, rows = _rows(kp)
    extra = kp.max_column + 1
    kp.cell(cx.HEADER_ROW, extra).value = "Done?"
    n = next(v for (code, _), v in rows.items() if code == "EMP-0281")
    kp.cell(n, extra).value = "Completed"
    res = _upload(client, rajesh, oct_, wb, dry_run=True).json()
    kpi_err = next(r for r in res["rows"] if r["sheet"] == cx.SHEET_KPIS and r["code"] == "EMP-0281")
    assert "not by a completion date" in " ".join(kpi_err["errors"])

    # Only the evaluator edits: Amit uploading Rajesh's file changes nothing (those people are outside Amit's scope).
    res = _upload(client, amit, oct_, wb, dry_run=False).json()
    assert res["applied"] == 0 and all("outside your scope" in " ".join(r["errors"]) for r in res["rows"])
    # The HOD sees everyone but is not their evaluator: refused per row with the person page's message.
    ci = wb[cx.SHEET_CHECKINS]
    head, rows = _rows(ci)
    ci.cell(rows[("EMP-0285", None)], head["Quality"]).value = 3
    res = _upload(client, auth("sales_hod"), oct_, wb, dry_run=True).json()
    hod_row = next(r for r in res["rows"] if r["sheet"] == cx.SHEET_CHECKINS and r["code"] == "EMP-0285")
    assert "evaluator" in " ".join(hod_row["errors"])

    # Closed months are untouched: the same kind of file cannot be uploaded to September.
    sep = months["Sep 2026"]
    assert sep["state"] == "Locked"
    sep_wb = _download(client, rajesh, sep)
    assert "closed" in sep_wb[cx.SHEET_CHECKINS]["A3"].value
    r = _upload(client, rajesh, sep, sep_wb, dry_run=True)
    assert r.status_code == 422 and "closed" in r.json()["detail"]["message"]
